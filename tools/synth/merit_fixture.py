"""Build tests/fixtures/merit_v7 from Merit's own synthetic outputs (ete-dashboard side_tools/merit/out/alteryx/csv).

Merit's synthetic data are fictitious, but its file register names real companies (payers, TPAs, the master cedant, a
reinsurer), and those names also sit inside LifeKey / ClaimKey. This repository keeps no client names (README), so
the fixture is Merit's output with exactly those values replaced, consistently in every file and key; every row and
every column is kept, so the fixture has Merit's real format.

    python tools/synth/merit_fixture.py <merit csv folder> [<out folder>]
Default out folder: tests/fixtures/merit_v7. Re-run when Merit's output contract changes; tests/test_merit.py compares
the committed fixture with a fresh build when an ete-dashboard clone sits next to this repository.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[2]
FILES = ("dim_life", "fact_member_month", "fact_claims", "run_manifest_claims")
KEY_COLS = {"LifeKey", "ClaimKey", "ClaimantKey"}          # '|'-separated keys: replaced segment by segment
NAMES = {                                                  # exact values only (never substrings of free text)
    "Orient Insurance PJSC": "Fictional Payer B",
    "United Fidelity Insurance Company": "Fictional Payer A",
    "UFIC": "FICT CEDANT",
    "NEXTCARE": "TPA ONE", "AAFIYA": "TPA TWO", "MEDNET": "TPA THREE", "NAS": "TPA FOUR",
    "GC-MRE": "RI-A", "GC-Optio": "RI-B",
}
HYPHEN = {"NEXTCARE": "T1", "AAFIYA": "T2", "MEDNET": "T3", "NAS": "T4", "UFIC": "FC"}   # inside ids, e.g. MC-NAS-1


def scrub_segment(s: str) -> str:
    if s in NAMES:
        return NAMES[s]
    return "-".join(HYPHEN.get(p, p) for p in s.split("-")) if "-" in s else s


def scrub_value(v, key: bool):
    """Exact value, or (key columns, and any value holding '|') each '|'-separated segment."""
    if not isinstance(v, str):
        return v
    if key or "|" in v:
        return "|".join(scrub_segment(s) for s in v.split("|"))
    return scrub_segment(v)


def build(src: Path, out: Path) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    counts = {}
    for f in FILES:
        df = pd.read_csv(src / f"{f}.csv", dtype=str, keep_default_na=False, encoding="utf-8-sig")
        for c in df.columns:
            df[c] = df[c].map(lambda v, k=(c in KEY_COLS): scrub_value(v, k))
        text = df.to_csv(index=False, lineterminator="\n")
        parts = set()
        for col in df.columns:
            for seg in df[col].str.split("|").explode():
                parts.add(seg)
                parts.update(seg.split("-"))
        left = [n for n in list(NAMES) + list(HYPHEN) if n in parts]
        if left:
            raise SystemExit(f"{f}: names left after scrubbing: {left}")
        (out / f"{f}.csv").write_text(text, encoding="utf-8")
        counts[f] = len(df)
    return counts


def test_control(src: Path, out: Path, payer: str = "") -> Path:
    """A Control workbook for the fixture: the synthetic Control with source_mode E, treaty years 2025-2026, quote year
    2027 (trend row added = 2026's) and the payer filter. Used by tests/test_merit.py and tools/check_merit.bat."""
    import shutil
    from openpyxl import load_workbook
    sys.path.insert(0, str(ROOT / "tools/excel"))
    import xlsx_norm
    shutil.copy(src, out)
    wb = load_workbook(out)
    ws = wb["Params"]
    vals = {"source_mode": "E", "quote_year": "2027", "years_included": "2025,2026", "treaty_start": "2026-01-01",
            "ete_payer_filter": payer or None, "ete_tpa_filter": None, "segment_field": "Policy Type"}
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value in vals:
            ws.cell(r, 3).value = vals[ws.cell(r, 1).value]
    tr = wb["Trend"]
    last = tr.max_row
    for c in range(1, tr.max_column + 1):
        tr.cell(last + 1, c).value = tr.cell(last, c).value
    tr.cell(last + 1, 1).value = 2027
    wb.save(out)
    return xlsx_norm.normalise(out)


if __name__ == "__main__":
    src = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "tests/fixtures/merit_v7"
    print(build(src, out))
