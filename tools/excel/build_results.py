"""Build the Results workbook template: excel/templates/XoL_Results.xlsx (docs/data-contracts.md §5; scope §7-8).

Presentation sheets (formulas only, no VBA): README, Inputs (display, return periods, expected-loss basis, loadings with
overrides), Pricing (indicative RI premium per option), Reinsurance Structures (Exhibit C1), Cedant Impact (Exhibit C2),
Expiring, Scenarios. Data sheets (out_*, scn_summary) carry headers only; stage 07 replaces them each run.

Formulas reach the data sheets only through workbook names such as ss_mean_recovery, each defined as
    INDEX(INDIRECT("'out_struct_summary'!$A$2:$BZ$60000"), 0, MATCH("mean_recovery", INDIRECT("'out_struct_summary'!$A$1:$BZ$1"), 0))
so a data sheet that Alteryx drops and re-creates never breaks a reference, and columns are found by header name.

    python tools/excel/build_results.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tools/excel"))
import results_spec as R  # noqa: E402
import xlsx_norm  # noqa: E402

OUT = ROOT / "excel/templates/XoL_Results.xlsx"
ABBR = {"out_params": "pa", "out_loadings": "lo", "out_run_log": "rl", "out_struct_summary": "ss",
        "out_struct_quantiles": "sq", "out_sim_summary": "ps", "out_ev_sim_summary": "es", "out_bc_summary": "bs",
        "out_ev_bc_summary": "eb", "out_exposed_lives": "el", "out_gross_summary": "gs", "out_expiring": "xp",
        "out_asif_by_year": "ai", "scn_summary": "sc"}
ROWS = 60000

F = "Arial"
TITLE = Font(name=F, size=14, bold=True, color="1F3864")
SUB = Font(name=F, size=10, italic=True, color="595959")
HEAD = Font(name=F, size=10, bold=True, color="FFFFFF")
BOLD = Font(name=F, size=10, bold=True)
BASE = Font(name=F, size=10)
INPUT = Font(name=F, size=10, color="0000FF")
SMALL = Font(name=F, size=8, color="808080")
HFILL = PatternFill("solid", fgColor="1F3864")
YFILL = PatternFill("solid", fgColor="FFF2CC")
GFILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
AMT = '#,##0;(#,##0);"-"'
AMT1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0'

# Pricing rows: PP options 1-4, EV options 1-4, SL options 1-4
OPTS = [("PP", o) for o in range(1, 5)] + [("EV", o) for o in range(1, 5)] + [("SL", o) for o in range(1, 5)]
P0 = 6                                    # first Pricing data row
PROW = {k: P0 + i for i, k in enumerate(OPTS)}
KEY = {"PP": "1 PP ", "EV": "3 EV ", "SL": "2 SL "}
STRUCTS = [("0 NONE", "Without RI")] + [(KEY[c] + str(o), f"{c} {o}") for c, o in OPTS] + [("4 PROGRAMME", "Programme")]


def name_for(sheet, col):
    return f"{ABBR[sheet]}_{col}"


def define_names(wb):
    for sheet, cols in R.SHEETS.items():
        if sheet not in ABBR:
            continue
        for c in cols:
            n = name_for(sheet, c)
            wb.defined_names[n] = DefinedName(
                n, attr_text=f'INDEX(INDIRECT("\'{sheet}\'!$A$2:$BZ${ROWS}"),0,'
                             f'MATCH("{c}",INDIRECT("\'{sheet}\'!$A$1:$BZ$1"),0))')


def cell(ws, ref, value, font=BASE, fmt=None, fill=None, align=None, comment=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if comment:
        c.comment = Comment(comment, "XoL tool")
    return c


def header(ws, row, labels, col0=1, wrap=True):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=t)
        c.font, c.fill = HEAD, HFILL
        c.alignment = Alignment(wrap_text=wrap, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 42


def widths(ws, w):
    for k, v in w.items():
        ws.column_dimensions[k].width = v


def title(ws, text, sub):
    cell(ws, "A1", text, TITLE)
    cell(ws, "A2", sub, SUB)
    ws.sheet_view.showGridLines = False


# ------------------------------------------------------------------------------------------------------------- Inputs
INP = {"ccy": "$E$5", "fx": "$E$6", "disp": "$E$7", "unit": "$E$8", "ref": "$E$9", "lives": "$E$10", "epi": "$E$11",
       "budget": "$E$12", "rp1": "$E$13", "rp2": "$E$14", "rp3": "$E$15", "basis": "$E$16", "factor": "$E$18",
       "cur": "$E$19"}
LOAD0 = 36                                 # used loadings rows 36-38 (PP, EV, SL), columns C..G, total H


def ii(k):
    return f"Inputs!{INP[k]}"


def rp_parse(k):
    t = "INDEX(pa_return_periods,1)"
    return f'=IFERROR(VALUE(TRIM(MID(SUBSTITUTE({t},",",REPT(" ",100)),{(k - 1) * 100 + 1},100))),{[5, 50, 250][k - 1]})'


def inputs(wb):
    ws = wb.create_sheet("Inputs")
    title(ws, "Inputs and loadings",
          "Yellow cells are optional overrides: blank uses the value from the run. Changing a loading or an override "
          "updates every exhibit without re-running Alteryx.")
    header(ws, 4, ["", "Item", "From the run", "Override", "Used", "Note"])
    items = [
        ("Quote currency", "=IFERROR(INDEX(pa_quote_ccy,1),\"\")", None, "Params quote_ccy"),
        ("1 unit of quote currency in USD", "=IFERROR(INDEX(pa_fx_quote_to_usd,1),1)", "0.000000", "Params fx_quote_to_usd"),
        ("Display currency (QUOTE / USD)", "=IFERROR(INDEX(pa_currency_display,1),\"QUOTE\")", None, "Exhibits only"),
        ("Display unit (1, 1000, 1000000)", 1000, NUM, "Amounts in the exhibits are divided by this"),
        ("Reference lives", "=IFERROR(INDEX(el_reference_lives,1),\"\")", NUM, "Premium per reference lives (C1)"),
        ("Lives in the quote year", "=IFERROR(INDEX(el_lives_quote_year,1),\"\")", NUM, "Stage 03 projection"),
        ("Expected premium income (EPI)", "=IFERROR(INDEX(pa_epi_quote_year,1),\"\")", AMT, "Params epi_quote_year"),
        ("Net claims budget", "=IFERROR(INDEX(gs_budget,1),\"\")", AMT,
         "Run value: net_claims_budget, else mean simulated gross. Stop-loss attachments were set in the run"),
        ("Return period 1 (years)", rp_parse(1), NUM, "Params return_periods"),
        ("Return period 2 (years)", rp_parse(2), NUM, ""),
        ("Return period 3 (years)", rp_parse(3), NUM, ""),
        ("Expected loss used for PP / EV", "BLENDED", None,
         "BLENDED (credibility blend), STOCHASTIC or DETERMINISTIC. Stop loss always uses the simulation"),
    ]
    for i, (lab, frm, fmt, note) in enumerate(items):
        r = 5 + i
        cell(ws, f"B{r}", lab)
        cell(ws, f"C{r}", frm, fmt=fmt, fill=GFILL)
        cell(ws, f"D{r}", None, INPUT, fmt=fmt, fill=YFILL)
        cell(ws, f"E{r}", f'=IF(D{r}="",C{r},D{r})', BOLD, fmt=fmt)
        cell(ws, f"F{r}", note, SUB)
    dv = DataValidation(type="list", formula1='"BLENDED,STOCHASTIC,DETERMINISTIC"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("D16")
    dv2 = DataValidation(type="list", formula1='"QUOTE,USD"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add("D7")
    cell(ws, "B18", "Display factor (currency / unit)")
    cell(ws, "E18", '=IF(E7="USD",E6,1)/E8', BOLD, fmt="0.000000")
    cell(ws, "B19", "Display currency label")
    cell(ws, "E19", '=IF(E7="USD","USD",E5)', BOLD)
    cell(ws, "B20", "Scenario of this run")
    cell(ws, "E20", '=IFERROR(INDEX(rl_value,MATCH("scenario_name",rl_item,0)),"")', BOLD)

    cols = ["Cover", "Reinsurer expenses", "Brokerage", "Profit margin", "Premium tax", "Risk load k x SD", "Total %"]
    lcol = ["reinsurer_expense_pct", "brokerage_pct", "profit_margin_pct", "premium_tax_pct", "risk_load_k"]
    cell(ws, "B22", "Loadings from the Control workbook", BOLD)
    header(ws, 23, cols, col0=2)
    cell(ws, "B28", "Overrides (blank = Control value)", BOLD)
    header(ws, 29, cols[:-1], col0=2)
    cell(ws, "B34", "Loadings used", BOLD)
    header(ws, 35, cols, col0=2)
    for j, cv in enumerate(("PP", "EV", "SL")):
        rc, ro, ru = 24 + j, 30 + j, LOAD0 + j
        for r_ in (rc, ro, ru):
            cell(ws, f"B{r_}", cv, BOLD)
        for k, lc in enumerate(lcol):
            col = L(3 + k)
            fmt = "0.00" if lc == "risk_load_k" else PCT
            cell(ws, f"{col}{rc}", f'=SUMIFS(lo_{lc},lo_cover,$B{rc})', fmt=fmt, fill=GFILL)
            cell(ws, f"{col}{ro}", None, INPUT, fmt=fmt, fill=YFILL)
            cell(ws, f"{col}{ru}", f'=IF({col}{ro}="",{col}{rc},{col}{ro})', BOLD, fmt=fmt)
        cell(ws, f"H{rc}", f"=SUM(C{rc}:F{rc})", fmt=PCT, fill=GFILL)
        cell(ws, f"H{ru}", f"=SUM(C{ru}:F{ru})", BOLD, fmt=PCT)
    cell(ws, "B40", "Premium = (expected loss + k x SD) / (1 - total %). With reinstatements, that is the expected total "
                    "premium; the upfront premium P solves P x (1 + expected reinstatement usage) = target.", SUB)
    widths(ws, {"A": 2, "B": 34, "C": 16, "D": 14, "E": 16, "F": 16, "G": 16, "H": 12})
    ws.column_dimensions["F"].width = 60
    return ws


# ------------------------------------------------------------------------------------------------------------ Pricing
PCOLS = ["Cover", "Option", "Structure key", "Present", "Label", "E[L] deterministic", "E[L] stochastic",
         "E[L] blended", "E[L] used", "SD of recoveries", "k", "Total loading %", "Target premium (incl. reinst.)",
         "E[reinstatement usage]", "Upfront annual premium", "E[reinstatement premium]", "Rate per life",
         "Rate % EPI", "Exposed lives", "Premium per exposed life", "Premium per reference lives",
         "Expected recoveries p.a.", "Expected cost of cover", "P(recoveries > premium)",
         "Recoveries exceed premium 1 in N yrs", "Upfront premium layer 1", "Upfront premium layer 2"]
PC = {n: L(i + 1) for i, n in enumerate(PCOLS)}


def pricing(wb):
    ws = wb.create_sheet("Pricing")
    title(ws, "Indicative reinsurance premium by option",
          "Quote currency, full units. Rows show only the options in the run. E[L] by layer is summed per option; "
          "the SD is that of the option's total annual recovery.")
    header(ws, 5, PCOLS)
    for (cv, o), r in PROW.items():
        c = lambda n: f"{PC[n]}{r}"  # noqa: E731
        sm = {"PP": "ps", "EV": "es"}.get(cv)
        lr = LOAD0 + ("PP", "EV", "SL").index(cv)
        cell(ws, c("Cover"), cv, BOLD)
        cell(ws, c("Option"), o)
        cell(ws, c("Structure key"), f'="{KEY[cv]}"&{c("Option")}', SMALL)
        cell(ws, c("Present"), f'=COUNTIF(ss_struct_key,{c("Structure key")})>0')
        cell(ws, c("Label"), f'=IF({c("Present")},INDEX(ss_l_label,MATCH({c("Structure key")},ss_struct_key,0)),"")')
        pres = c("Present")
        if sm:
            det = f"SUMIFS({sm}_expected_cost_det,{sm}_option_id,{c('Option')})"
            sto = f"SUMIFS({sm}_mean_recovery,{sm}_option_id,{c('Option')})"
            bl = f"SUMIFS({sm}_expected_loss_blended,{sm}_option_id,{c('Option')})"
            ru = (f"IFERROR(SUMPRODUCT(({sm}_option_id={c('Option')})*{sm}_mean_recovery*{sm}_mean_reinst_usage)"
                  f"/{sto},0)")
            l1 = f"IFERROR({c('Upfront annual premium')}*SUMIFS({sm}_mean_recovery,{sm}_option_id,{c('Option')},{sm}_layer,1)/{sto},{c('Upfront annual premium')})"
        else:
            sto = f"SUMIFS(ss_mean_recovery,ss_struct_key,{c('Structure key')})"
            det, bl, ru, l1 = '""', sto, "0", c("Upfront annual premium")
        cell(ws, c("E[L] deterministic"), f"=IF({pres},{det},\"\")", fmt=AMT)
        cell(ws, c("E[L] stochastic"), f"=IF({pres},{sto},\"\")", fmt=AMT)
        cell(ws, c("E[L] blended"), f"=IF({pres},{bl},\"\")", fmt=AMT)
        used = (c("E[L] stochastic") if cv == "SL" else
                f'IF({ii("basis")}="DETERMINISTIC",{c("E[L] deterministic")},IF({ii("basis")}="STOCHASTIC",'
                f'{c("E[L] stochastic")},{c("E[L] blended")}))')
        cell(ws, c("E[L] used"), f"=IF({pres},{used},\"\")", BOLD, fmt=AMT)
        cell(ws, c("SD of recoveries"), f"=IF({pres},SUMIFS(ss_sd_recovery,ss_struct_key,{c('Structure key')}),\"\")",
             fmt=AMT)
        cell(ws, c("k"), f"=Inputs!$G${lr}", fmt="0.00")
        cell(ws, c("Total loading %"), f"=Inputs!$H${lr}", fmt=PCT)
        cell(ws, c("Target premium (incl. reinst.)"),
             f"=IF({pres},({c('E[L] used')}+{c('k')}*{c('SD of recoveries')})/(1-{c('Total loading %')}),\"\")", fmt=AMT)
        cell(ws, c("E[reinstatement usage]"), f"=IF({pres},{ru},\"\")", fmt="0.000")
        cell(ws, c("Upfront annual premium"),
             f"=IF({pres},{c('Target premium (incl. reinst.)')}/(1+{c('E[reinstatement usage]')}),\"\")", BOLD, fmt=AMT)
        cell(ws, c("E[reinstatement premium]"),
             f"=IF({pres},{c('Upfront annual premium')}*{c('E[reinstatement usage]')},\"\")", fmt=AMT)
        cell(ws, c("Rate per life"), f"=IF({pres},{c('Upfront annual premium')}/{ii('lives')},\"\")", fmt="#,##0.00")
        cell(ws, c("Rate % EPI"), f"=IF({pres},{c('Upfront annual premium')}/{ii('epi')},\"\")", fmt="0.000%")
        cell(ws, c("Exposed lives"),
             f"=IF({pres},SUMIFS(el_exposed_lives,el_cover,{c('Cover')},el_option_id,{c('Option')},el_layer,1),\"\")",
             fmt=NUM)
        cell(ws, c("Premium per exposed life"),
             f"=IF({pres},IFERROR({c('Upfront annual premium')}/{c('Exposed lives')},0),\"\")", fmt="#,##0.00")
        cell(ws, c("Premium per reference lives"),
             f"=IF({pres},{c('Premium per exposed life')}*{ii('ref')},\"\")", fmt=AMT)
        cell(ws, c("Expected recoveries p.a."), f"=IF({pres},{c('E[L] used')},\"\")", fmt=AMT)
        cell(ws, c("Expected cost of cover"),
             f"=IF({pres},{c('Upfront annual premium')}+{c('E[reinstatement premium]')}-{c('E[L] used')},\"\")", fmt=AMT)
        cell(ws, c("P(recoveries > premium)"),
             f"=IF({pres},COUNTIFS(sq_struct_key,{c('Structure key')},sq_recovery,\">\"&{c('Upfront annual premium')})"
             f"/MAX(1,COUNTIF(sq_struct_key,{c('Structure key')})),\"\")", fmt="0.0%")
        cell(ws, c("Recoveries exceed premium 1 in N yrs"),
             f"=IF({pres},IF({c('P(recoveries > premium)')}>0,1/{c('P(recoveries > premium)')},\"> 999\"),\"\")",
             BOLD, fmt="0.0")
        cell(ws, c("Upfront premium layer 1"), f"=IF({pres},{l1},\"\")", fmt=AMT)
        cell(ws, c("Upfront premium layer 2"),
             f"=IF({pres},{c('Upfront annual premium')}-{c('Upfront premium layer 1')},\"\")", fmt=AMT)
    cell(ws, f"A{P0 + len(OPTS) + 1}",
         "Premium is split between layers in proportion to their simulated expected recoveries (used for the "
         "reinstatement premium of each layer). 1 in N = 1 / share of the 999 simulated recovery quantiles above the "
         "premium.", SUB)
    widths(ws, {L(i + 1): 14 for i in range(len(PCOLS))})
    ws.column_dimensions["E"].width = 26
    ws.freeze_panes = "F6"
    return ws


def p(col, cv, o):
    return f"Pricing!${PC[col]}${PROW[(cv, o)]}"


# --------------------------------------------------------------------------------------- Reinsurance Structures (C1)
def structures(wb):
    ws = wb.create_sheet("Reinsurance Structures")
    title(ws, "Reinsurance structures & options", "")
    cell(ws, "A2", f'="Amounts in "&{ii("cur")}&IF({ii("unit")}=1000," thousands",IF({ii("unit")}=1000000," millions",""))'
                   f'&". Indicative premiums from high-level loadings (Inputs)."', SUB)
    cols = ["Option", "Cover", "Exposed lives", "Annual premium", "Premium per exposed life p.a.",
            "Premium per reference lives", "Expected recoveries p.a.", "Expected cost of cover",
            "Recoveries exceed premium 1 in N yrs"]
    header(ws, 4, cols, col0=2)
    cell(ws, "G4", f'="Premium per "&TEXT({ii("ref")},"#,##0")&" lives"', HEAD,
         align=Alignment(wrap_text=True, horizontal="center", vertical="center"), fill=HFILL)
    fac, cur = ii("factor"), ii("cur")
    r = 5
    for cv, o in OPTS:
        pres = p("Present", cv, o)
        name = {"PP": "Per-person XoL ", "EV": "Per-event XoL ", "SL": "Stop loss "}[cv]
        cell(ws, f"B{r}", f'=IF({pres},"{name}"&{o},"")', BOLD)
        if cv == "SL":
            k = f'"2 SL "&{o}'
            g = lambda col: f"SUMIFS(ss_{col},ss_struct_key,{k})"  # noqa: E731
            m = f'INDEX(ss_sl_measure,MATCH({k},ss_struct_key,0))'
            fxm = f'IF({ii("disp")}="USD",{ii("fx")},1)/1000000'
            amt = (f'{cur}&" "&TEXT({g("limit_amount")}*{fxm},"#,##0.0")&"m xs "&{cur}&" "&'
                   f'TEXT({g("ap_amount")}*{fxm},"#,##0.0")&"m"')
            desc = (f'IF({pres},IFERROR(IF({m}="PCT_BUDGET",TEXT({g("limit_pct_budget")},"0%")&" xs "&'
                    f'TEXT({g("ap_pct_budget")},"0%")&" of budget ("&{amt}&")",IF({m}="PCT_EPI",TEXT({g("limit_pct_epi")},"0%")'
                    f'&" xs "&TEXT({g("ap_pct_epi")},"0%")&" of EPI ("&{amt}&")",{amt}))&IF(INDEX(ss_pp_option,MATCH({k},'
                    f'ss_struct_key,0))>0," net of XoL",""),{p("Label", cv, o)}),"")')
            cell(ws, f"C{r}", "=" + desc)
        else:
            cell(ws, f"C{r}", f"={p('Label', cv, o)}")
        cell(ws, f"D{r}", f'=IF({pres},{p("Exposed lives", cv, o)},"")', fmt=NUM)
        for col, src in (("E", "Upfront annual premium"), ("F", "Premium per exposed life"),
                         ("G", "Premium per reference lives"), ("H", "Expected recoveries p.a."),
                         ("I", "Expected cost of cover")):
            fmt = "#,##0.00" if col == "F" else AMT1
            mult = f"*IF({ii('disp')}=\"USD\",{ii('fx')},1)" if col == "F" else f"*{fac}"
            cell(ws, f"{col}{r}", f'=IF({pres},{p(src, cv, o)}{mult},"")', fmt=fmt)
        cell(ws, f"J{r}", f'=IF({pres},{p("Recoveries exceed premium 1 in N yrs", cv, o)},"")', BOLD, fmt="0.0",
             align=Alignment(horizontal="center"))
        for col in "BCDEFGHIJ":
            ws[f"{col}{r}"].border = THIN
        r += 1
    cell(ws, f"B{r + 1}", "Premium per exposed life is in the display currency, full units. Per-person exposed lives "
                          "are lives whose benefit limit exceeds the attachment point; stop loss and per-event use all "
                          "lives.", SUB)
    widths(ws, {"A": 2, "B": 22, "C": 46, "D": 12, "E": 14, "F": 14, "G": 14, "H": 14, "I": 14, "J": 14})
    return ws


# ------------------------------------------------------------------------------------------------ Cedant Impact (C2)
def impact(wb):
    ws = wb.create_sheet("Cedant Impact")
    title(ws, "Cedant impact: surplus / (loss) against the net claims budget", "")
    cell(ws, "A2", f'="Amounts in "&{ii("cur")}&IF({ii("unit")}=1000," thousands",IF({ii("unit")}=1000000," millions",""))'
                   f'&". Surplus = budget - (claims net of recoveries + RI premium incl. reinstatement premium); '
                   f'return-period values are quantiles of the joint simulated outcome."', SUB)
    fac = ii("factor")
    c0 = 3
    cols = [L(c0 + i) for i in range(len(STRUCTS))]
    cell(ws, "B5", "Structure key", SMALL)
    cell(ws, "B6", "Structure", BOLD)
    cell(ws, "B7", "Annual RI premium (incl. expected reinstatement premium)")
    cell(ws, "B8", "Net claims budget")
    cell(ws, "B10", "Surplus / (loss) against the net claims budget", BOLD)
    cell(ws, "B11", "Average year")
    for k in range(3):
        cell(ws, f"B{12 + k}", f'="1 in "&TEXT({ii("rp" + str(k + 1))},"0")&" years"')
    cell(ws, "B17", "Workings (quote currency, full units)", BOLD)
    labs = ["PP option", "EV option", "SL option", "Premium PP layer 1", "Premium PP layer 2", "Premium EV layer 1",
            "Premium EV layer 2", "Premium SL", "Upfront premium", "E[reinstatement premium]", "Mean net claims"]
    for i, t in enumerate(labs):
        cell(ws, f"B{18 + i}", t, SMALL)
    for k in range(3):
        cell(ws, f"B{30 + k}", f'="q (per mille), 1 in "&TEXT({ii("rp" + str(k + 1))},"0")', SMALL)
        cell(ws, f"A{30 + k}", f'=ROUND(1000*(1-1/{ii("rp" + str(k + 1))}),0)', SMALL)
        cell(ws, f"B{33 + k}", f'="Net claims at 1 in "&TEXT({ii("rp" + str(k + 1))},"0")', SMALL)
        cell(ws, f"B{36 + k}", f'="Reinstatement premium at 1 in "&TEXT({ii("rp" + str(k + 1))},"0")', SMALL)
    for (key, short), col in zip(STRUCTS, cols):
        x = lambda r: f"{col}{r}"  # noqa: E731
        k = f"{col}$5"
        cell(ws, x(5), key, SMALL)
        pres = f"COUNTIF(ss_struct_key,{k})>0"
        cell(ws, x(6), f'=IF({pres},INDEX(ss_l_label,MATCH({k},ss_struct_key,0))&IF(AND(LEFT({k},1)="2",{x(18)}+{x(19)}>0),'
                       f'" + XoL","") ,"")', BOLD, align=Alignment(wrap_text=True, horizontal="center"))
        for r_, nm in ((18, "pp_option"), (19, "ev_option"), (20, "sl_option")):
            cell(ws, x(r_), f"=IFERROR(INDEX(ss_{nm},MATCH({k},ss_struct_key,0)),0)", SMALL)
        pr = lambda colname, cv, rr: (f"IF({rr}>0,INDEX(Pricing!${PC[colname]}${PROW[(cv, 1)]}:${PC[colname]}"  # noqa: E731
                                      f"${PROW[(cv, 4)]},{rr}),0)")
        cell(ws, x(21), "=" + pr("Upfront premium layer 1", "PP", x(18)), SMALL, fmt=AMT)
        cell(ws, x(22), "=" + pr("Upfront premium layer 2", "PP", x(18)), SMALL, fmt=AMT)
        cell(ws, x(23), "=" + pr("Upfront premium layer 1", "EV", x(19)), SMALL, fmt=AMT)
        cell(ws, x(24), "=" + pr("Upfront premium layer 2", "EV", x(19)), SMALL, fmt=AMT)
        cell(ws, x(25), "=" + pr("Upfront annual premium", "SL", x(20)), SMALL, fmt=AMT)
        cell(ws, x(26), f"=SUM({x(21)}:{x(25)})", SMALL, fmt=AMT)
        mr = lambda f: f"SUMIFS(ss_{f},ss_struct_key,{k})"  # noqa: E731
        cell(ws, x(27), f"={x(21)}*{mr('mean_ru_l1')}+{x(22)}*{mr('mean_ru_l2')}+{x(23)}*{mr('mean_ru_ev1')}"
                        f"+{x(24)}*{mr('mean_ru_ev2')}", SMALL, fmt=AMT)
        cell(ws, x(28), f"={mr('mean_net')}", SMALL, fmt=AMT)
        for kk in range(3):
            qm = f"$A${30 + kk}"
            q = lambda f: f"SUMIFS(sq_{f},sq_struct_key,{k},sq_qm,{qm})"  # noqa: E731
            cell(ws, x(33 + kk), f"={q('net_claims')}", SMALL, fmt=AMT)
            cell(ws, x(36 + kk), f"={x(21)}*{q('ru_l1_at_q')}+{x(22)}*{q('ru_l2_at_q')}+{x(23)}*{q('ru_ev1_at_q')}"
                                 f"+{x(24)}*{q('ru_ev2_at_q')}", SMALL, fmt=AMT)
        cell(ws, x(7), f'=IF({pres},({x(26)}+{x(27)})*{fac},"")', fmt=AMT1)
        cell(ws, x(8), f'=IF({pres},{ii("budget")}*{fac},"")', fmt=AMT1)
        cell(ws, x(11), f'=IF({pres},({ii("budget")}-{x(28)}-{x(26)}-{x(27)})*{fac},"")', BOLD, fmt=AMT1)
        for kk in range(3):
            cell(ws, x(12 + kk), f'=IF({pres},({ii("budget")}-{x(33 + kk)}-{x(26)}-{x(36 + kk)})*{fac},"")', BOLD,
                 fmt=AMT1)
        for r_ in (6, 7, 8, 11, 12, 13, 14):
            ws[x(r_)].border = THIN
        ws.column_dimensions[col].width = 14
    widths(ws, {"A": 6, "B": 48})
    ws.row_dimensions[6].height = 45
    ws.freeze_panes = "C7"
    return ws


# --------------------------------------------------------------------------------------------------------- Expiring
def expiring(wb):
    ws = wb.create_sheet("Expiring")
    title(ws, "Expiring structure: rate comparison and as-if experience",
          "Expiring rate applied to the quote year, then trended by the matched option's leveraged trend "
          "(base trend if there is no match); indicated = upfront premium of the option with the same AP and limit.")
    cols = ["Cover", "Layer", "AP", "Limit", "Rate basis", "Rate", "Matched option", "Leveraged trend",
            "Expiring premium (quote year)", "Expiring premium trended", "Indicated premium", "Change"]
    header(ws, 4, cols, col0=1)
    xcol = ["cover", "layer", "ap", "limit", "rate_basis", "rate", "matched_option", "lev_trend_used",
            "expiring_premium_q", "expiring_premium_trended"]
    fm = [None, "0", AMT, AMT, None, "#,##0.00##", "0", PCT, AMT, AMT]
    for i in range(4):
        r = 5 + i
        for j, (c_, f_) in enumerate(zip(xcol, fm)):
            cell(ws, f"{L(j + 1)}{r}", f'=IFERROR(INDEX(xp_{c_},{i + 1})&"","")' if f_ is None else
                 f'=IFERROR(IF(INDEX(xp_{c_},{i + 1})="","",INDEX(xp_{c_},{i + 1})),"")', fmt=f_)
        cell(ws, f"K{r}", f'=IFERROR(IF(AND(A{r}="PP",G{r}<>""),INDEX(Pricing!${PC["Upfront annual premium"]}$'
                          f'{PROW[("PP", 1)]}:${PC["Upfront annual premium"]}${PROW[("PP", 4)]},G{r}),""),"")', BOLD, fmt=AMT)
        cell(ws, f"L{r}", f'=IFERROR(K{r}/J{r}-1,"")', fmt=PCT)
    cell(ws, "A11", "As-if experience of the matched option: layer loss trended to the quote year against the premium "
                    "at the expiring rate on that year's exposure, also at the quote-year level; actual (historical) "
                    "recoveries against the same premium at the year's own level", BOLD)
    cols2 = ["Treaty year", "Matched option", "Exposure (life-years)", "As-if recoveries (quote-year level)",
             "As-if premium (quote-year level)", "As-if loss ratio", "Premium at the year's level", "Actual recoveries",
             "Actual loss ratio"]
    header(ws, 12, cols2, col0=1)
    acol = ["ty", "matched_option", "exposure_ly_full", "asif_recovery", "asif_premium"]
    afm = ["0", "0", NUM, AMT, AMT]
    for i in range(12):
        r = 13 + i
        for j, (c_, f_) in enumerate(zip(acol, afm)):
            cell(ws, f"{L(j + 1)}{r}", f'=IFERROR(IF(INDEX(ai_{c_},{i + 1})="","",INDEX(ai_{c_},{i + 1})),"")', fmt=f_)
        cell(ws, f"F{r}", f'=IFERROR(D{r}/E{r},"")', fmt=PCT)
        for col, c_ in (("G", "premium_year_level"), ("H", "actual_recovery")):
            cell(ws, f"{col}{r}", f'=IFERROR(IF(INDEX(ai_{c_},{i + 1})="","",INDEX(ai_{c_},{i + 1})),"")', fmt=AMT)
        cell(ws, f"I{r}", f'=IFERROR(H{r}/G{r},"")', fmt=PCT)
    widths(ws, {L(i + 1): 15 for i in range(12)})
    return ws


# -------------------------------------------------------------------------------------------------------- Scenarios
def scenarios(wb):
    ws = wb.create_sheet("Scenarios")
    title(ws, "Sensitivity scenarios", "Each run with Params scenario_id adds its rows to scn_summary (a re-run of the "
                                       "same id replaces them). Pick a structure key below.")
    cell(ws, "B4", "Structure key", BOLD)
    cell(ws, "C4", "1 PP 1", INPUT, fill=YFILL)
    dv = DataValidation(type="list", formula1='"' + ",".join(k for k, _ in STRUCTS) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("C4")
    cols = ["Scenario id", "Scenario", "Run at", "Net claims budget", "Mean recovery", "SD of recoveries",
            "Mean net claims", "Expected loss (blended)"]
    header(ws, 6, cols, col0=1)
    for i in range(10):
        r = 7 + i
        cell(ws, f"A{r}", i)
        has = f'COUNTIFS(sc_scenario_id,$A{r},sc_struct_key,$C$4)>0'
        cell(ws, f"B{r}", f'=IFERROR(INDEX(sc_scenario_name,MATCH($A{r},sc_scenario_id,0)),"")')
        cell(ws, f"C{r}", f'=IFERROR(INDEX(sc_run_at,MATCH($A{r},sc_scenario_id,0)),"")')
        for col, f_ in (("D", "budget"), ("E", "mean_recovery"), ("F", "sd_recovery"), ("G", "mean_net"),
                        ("H", "expected_loss_blended")):
            cell(ws, f"{col}{r}", f'=IF({has},SUMIFS(sc_{f_},sc_scenario_id,$A{r},sc_struct_key,$C$4),"")', fmt=AMT)
    widths(ws, {"A": 12, "B": 30, "C": 18, "D": 16, "E": 16, "F": 16, "G": 16, "H": 18})
    return ws


# ----------------------------------------------------------------------------------------------------------- README
def readme(wb):
    ws = wb.active
    ws.title = "README"
    title(ws, "XoL pricing tool - Results workbook", "Broker view: expected claims cost, indicative RI cost, cedant impact")
    lines = [
        ("How it is filled", BOLD),
        ("1. Copy this template once to data/output/<live|synth>/XoL_Results.xlsx (next to the stage folders).", BASE),
        ("2. Run the workflows 00 -> 01 -> 02 -> 03 -> 04 -> 05 -> 06 -> 07 in Alteryx. Stage 07 replaces the data "
         "sheets (out_*, scn_summary); close this workbook in Excel before running it.", BASE),
        ("3. Open the workbook. Every exhibit recalculates from the data sheets and the Inputs sheet.", BASE),
        ("", BASE),
        ("Sheets", BOLD),
        ("Inputs - display currency and unit, reference lives, return periods, expected-loss basis, loadings "
         "(Control values with optional overrides).", BASE),
        ("Pricing - per option: expected loss (deterministic, stochastic, credibility-blended), SD, target and upfront "
         "premium, reinstatement premium, rates, exposed lives, 1-in-N recoveries exceed premium.", BASE),
        ("Reinsurance Structures - Exhibit C1 (client format).", BASE),
        ("Cedant Impact - Exhibit C2: surplus / (loss) against the net claims budget for the average year and the "
         "return periods, without RI and with each structure.", BASE),
        ("Expiring - expiring rate trended vs indicated premium; as-if experience by treaty year.", BASE),
        ("Scenarios - one row per sensitivity run (Params scenario_id).", BASE),
        ("out_* - data written by Alteryx (do not edit; they are replaced each run). out_run_log lists the "
         "parameters, justifications, scenario and data mode of the run; out_warnings every check (STOP / FLAG / INFO).",
         BASE),
        ("", BASE),
        ("Conventions", BOLD),
        ("Yellow cells are inputs (blue text). Grey cells show the run value. Formulas read the data sheets through "
         "names such as ss_mean_recovery (Formulas > Name Manager), defined with INDIRECT so that a data sheet "
         "replaced by Alteryx never breaks a reference.", BASE),
        ("No VBA, no Power Query: every calculation is a cell formula.", BASE),
        ("Caveats: per-person completion factors from portfolio development may understate development above an "
         "attachment point; per-event frequency rests on few events (see credibility); the stop-loss attritional CoV "
         "rests on 2-5 annual points (floor applies). Risk transfer and capital are out of scope.", BASE),
    ]
    for i, (t, f_) in enumerate(lines):
        c = cell(ws, f"A{4 + i}", t, f_)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 130
    return ws


def data_sheets(wb):
    for sheet, cols in R.SHEETS.items():
        ws = wb.create_sheet(sheet)
        ws.sheet_properties.tabColor = "A6A6A6"
        for i, c in enumerate(cols, 1):
            ws.cell(row=1, column=i, value=c).font = BOLD


def build(path=OUT):
    wb = Workbook()
    readme(wb)
    inputs(wb)
    pricing(wb)
    structures(wb)
    impact(wb)
    expiring(wb)
    scenarios(wb)
    data_sheets(wb)
    define_names(wb)
    wb.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    xlsx_norm.normalise(path)
    return path



if __name__ == "__main__":
    print("wrote", build())
