"""Build the Control workbook: the blank template and the synthetic instance.

    python tools/excel/build_control.py
        -> excel/templates/XoL_Control.xlsx           (defaults only; client fields blank)
        -> data/synthetic/XoL_Control_synthetic.xlsx  (filled for the synthetic portfolio)
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import control_spec as S  # noqa: E402
import xlsx_norm  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
NAVY, INPUT, NOTE = "1F3864", "FFF2CC", "F2F2F2"


def _style_sheet(ws, title: str, note: str, cols: list) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, color=NAVY)
    ws["A2"] = note
    ws["A2"].font = Font(italic=True, size=9, color="595959")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(cols), 6))
    ws.row_dimensions[2].height = 48
    for j, (name, *_rest) in enumerate(cols, 1):
        c = ws.cell(row=S.HEADER_ROW, column=j, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        ws.column_dimensions[get_column_letter(j)].width = max(14, min(42, len(name) + 4))
    ws.freeze_panes = ws.cell(row=S.HEADER_ROW + 1, column=1)


def _write_rows(ws, cols: list, rows: list[dict], n_blank: int = 0) -> None:
    for i, r in enumerate(rows + [{}] * n_blank):
        for j, (name, kind, allowed, _req) in enumerate(cols, 1):
            v = r.get(name)
            c = ws.cell(row=S.HEADER_ROW + 1 + i, column=j, value=v)
            if name in ("description", "allowed_values", "group", "unit", "param_id", "indicated_factor"):
                c.fill = PatternFill("solid", fgColor=NOTE)
            else:
                c.fill = PatternFill("solid", fgColor=INPUT)
            if kind == "NUMBER" and name.endswith("pct"):
                c.number_format = "0.00%"
            if kind == "TEXT" or ws.title == "Params":
                c.number_format = "@"
    last = S.HEADER_ROW + len(rows) + n_blank
    for j, (name, kind, allowed, _req) in enumerate(cols, 1):
        if allowed and kind in ("LIST", "INT") and "," in allowed:
            dv = DataValidation(type="list", formula1=f'"{allowed}"', allow_blank=True)
            ws.add_data_validation(dv)
            col = get_column_letter(j)
            dv.add(f"{col}{S.HEADER_ROW + 1}:{col}{max(last, S.HEADER_ROW + 50)}")


def _params_sheet(ws, values: dict) -> None:
    cols = S.TABLES["Params"][2]
    rows = []
    for pid, group, unit, kind, allowed, req, desc in S.PARAMS:
        rows.append({"param_id": pid, "group": group, "value": values.get(pid, S.DEFAULTS.get(pid)),
                     "unit": unit, "allowed_values": allowed, "description": desc + (" [required]" if req else ""),
                     "justification": values.get(f"just:{pid}")})
    _write_rows(ws, cols, rows)
    for i, (pid, *_r) in enumerate(S.PARAMS):
        allowed = S.PARAM_ALLOWED.get(pid)
        if allowed:
            dv = DataValidation(type="list", formula1=f'"{allowed}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"C{S.HEADER_ROW + 1 + i}")
    ws.column_dimensions["F"].width = 70
    ws.column_dimensions["G"].width = 50


def _readme(wb, synthetic: bool) -> None:
    ws = wb.create_sheet("README", 0)
    lines = [
        ("XoL Pricing Tool - Control workbook", True),
        (f"Version {S.VERSION}. {'SYNTHETIC INSTANCE - fictional portfolio for testing.' if synthetic else 'Template.'}", False),
        ("", False),
        ("Alteryx reads every sheet below by name; the header is in row %d. Do not rename sheets or headers." % S.HEADER_ROW, False),
        ("Yellow cells are inputs. Grey cells are labels or written back by Alteryx.", False),
        ("Every judgement input (trend, completion, weights, incidence) needs a justification - it goes into the run log.", False),
        ("", False),
    ] + [(f"{s}: {S.TABLES[s][0]}", False) for s in S.SHEETS]
    for i, (t, b) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(bold=b, size=13 if b else 10, color=NAVY if b else "000000")
    ws.column_dimensions["A"].width = 110


def build(path: Path, values: dict, tables: dict, synthetic: bool) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in S.SHEETS:
        title, note, cols = S.TABLES[sheet]
        ws = wb.create_sheet(sheet)
        _style_sheet(ws, title, note, cols)
        if sheet == "Params":
            _params_sheet(ws, values)
        else:
            _write_rows(ws, cols, tables.get(sheet, []), n_blank=0 if tables.get(sheet) else 5)
    _readme(wb, synthetic)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    xlsx_norm.normalise(path)


# --------------------------------------------------------------------------------------------- content
def default_tables() -> dict:
    lo = [{"cover": c, "reinsurer_expense_pct": 0.05, "brokerage_pct": 0.10, "profit_margin_pct": 0.10,
           "premium_tax_pct": 0.0, "risk_load_k": 0.0} for c in ("PP", "EV", "SL")]
    return {
        "Loadings": lo,
        "EventRule": [{"id_source": "CLUSTER", "window_days": 3, "group_key": "SCHEME+SITE", "min_lives": 3}],
        "Programme": [{"step": 1, "cover": "QS", "active": 0}, {"step": 2, "cover": "PP", "option_id": 1, "active": 1},
                      {"step": 3, "cover": "EV", "option_id": 1, "active": 1},
                      {"step": 4, "cover": "SL", "option_id": 1, "active": 1}],
    }


def synthetic_content() -> tuple[dict, dict]:
    truth = json.loads((ROOT / "tests/reconciliation/synthetic_truth.json").read_text())
    m = pd.read_csv(ROOT / "data/synthetic/line/members.csv", parse_dates=["policy_start"])
    m["aed"] = m.written_premium.where(m.currency == "AED", m.written_premium * 3.6725)
    w = m.groupby(m.policy_start.dt.year).agg(prem=("aed", "sum"), lives=("member_id", "nunique"))
    g = w.lives[2024] / w.lives[2023]
    epi = float(w.prem[2024] * (1.07 * g) ** 2)

    years = [2021, 2022, 2023, 2024, 2025]
    ra = truth["per_person"]["RA"]
    comp = []

    def ratio(a, b, key):
        x, y = ra[a].get(str(yr), {}).get(key, 0), ra[b].get(str(yr), {}).get(key, 0)
        return round(x / y, 4) if y else 1.0
    for yr in years:
        steps = {"IBNR": ("earned_ultimate", "reported_at_valuation_true_amount"),
                 "UNEARNED": ("written_cohorts_ultimate", "earned_ultimate"),
                 "UNWRITTEN": ("ultimate_all_cohorts", "written_cohorts_ultimate")}
        for comp_name, (a, b) in steps.items():
            sides = {"PRE_AP": 1.0, "POST_AP_AMOUNT": ratio(a, b, "L100xs_amount"),
                     "POST_AP_COUNT": ratio(a, b, "L100xs_count")}
            if comp_name == "UNWRITTEN":
                sides["PRE_AP"] = 1.0
            sides["ATTRITIONAL"] = ratio(a, b, "total")
            for side, f in sides.items():
                basis = "total claims" if side == "ATTRITIONAL" else "reference layer 400k xs 100k"
                comp.append({"year": yr, "component": comp_name, "side": side, "ap_band_from": 0,
                             "selected_factor": f,
                             "justification": f"Synthetic: set from synthetic_truth.json, {basis}" if f != 1.0 else None})

    values = {
        "cedant": "Synthetic Insurance Co (fictional)", "country": "UAE", "line": "Group medical",
        "quote_ccy": "AED", "fx_quote_to_usd": "0.272294", "treaty_start": "2026-01-01", "quote_year": "2026",
        "run_label": "synthetic-test", "epi_quote_year": f"{epi:.0f}", "years_included": "2021,2022,2023,2024,2025",
        "reference_lives": "", "just:incidence_adj": "No explained trend outside the trend assumption",
    }
    t = default_tables()
    t.update({
        "FX": [{"currency": "AED", "priced_rate_to_quote_ccy": 1.0, "rate_date": "2025-09-30"},
               {"currency": "USD", "priced_rate_to_quote_ccy": 3.6725, "rate_date": "2025-09-30"}],
        "Limits": [{"plan": "A", "annual_limit": 1_000_000}, {"plan": "B", "annual_limit": 500_000},
                   {"plan": "C", "annual_limit": 150_000}],
        "Trend": [{"year": y, "client_pct": 0, "market_pct": 0, "inflation_pct": 0, "residual_pct": 0.08,
                   "weight_client": 0, "weight_market": 0, "weight_inflation": 0, "one_off_adj": 0,
                   "premium_trend_pct": 0.07} for y in range(2020, 2027)],
        "Completion": comp,
        "Weights": [{"year": y, "months_complete": 9 if y == 2025 else 12, "use_default": 1} for y in years],
        "Layers": [
            {"cover": "PP", "option_id": 1, "layer": 1, "label": "400k xs 100k", "limit_basis": "AMOUNT",
             "ap": 100_000, "limit": 400_000, "active": 1},
            {"cover": "PP", "option_id": 2, "layer": 1, "label": "750k xs 250k", "limit_basis": "AMOUNT",
             "ap": 250_000, "limit": 750_000, "active": 1},
            {"cover": "PP", "option_id": 3, "layer": 1, "label": "xs 75k to benefit limit",
             "limit_basis": "TO_BENEFIT_LIMIT", "ap": 75_000, "active": 1},
            {"cover": "PP", "option_id": 4, "layer": 1, "label": "xs 150k to benefit limit",
             "limit_basis": "TO_BENEFIT_LIMIT", "ap": 150_000, "active": 1},
            {"cover": "EV", "option_id": 1, "layer": 1, "label": "1m xs 250k per event", "limit_basis": "AMOUNT",
             "ap": 250_000, "limit": 1_000_000, "reinstatements_n": 1, "reinstatement_cost_pct": 1.0,
             "reinst_pro_rata_time": 0, "active": 1},
        ] + [{"cover": "SL", "option_id": i + 1, "layer": 1, "label": f"{int(l*100)}% xs {int(a*100)}% of budget",
              "sl_claims_basis": "GROSS", "sl_measure": "PCT_BUDGET", "sl_ap": a, "sl_limit": l, "active": 1}
             for i, (l, a) in enumerate([(0.10, 1.00), (0.20, 1.00), (0.10, 1.10), (0.20, 1.10)])],
        "EvScenarios": [{"scenario_name": "Worksite transport accident", "lives": 10, "cost_per_life": 150_000,
                         "return_period_yrs": 50}],
        "Expiring": [{"cover": "PP", "layer": 1, "ap": 100_000, "limit": 400_000, "rate_basis": "PPPY", "rate": 950,
                      "recoveries_by_year": "2021:2650000; 2022:3480000; 2023:3610000; 2024:2900000"}],
        "Scenarios": [
            {"scenario_id": 1, "scenario_name": "Trend +2pp", "target_table": "Trend", "target_key": "ALL",
             "target_column": "residual_pct", "operation": "ADD", "override_value": "0.02"},
            {"scenario_id": 2, "scenario_name": "IBNR post-AP +10%", "target_table": "Completion",
             "target_key": "component=IBNR", "target_column": "selected_factor", "operation": "MULT",
             "override_value": "1.1"},
        ],
    })
    return values, t


def main() -> None:
    build(ROOT / "excel/templates/XoL_Control.xlsx", {}, default_tables(), synthetic=False)
    v, t = synthetic_content()
    build(ROOT / "data/synthetic/XoL_Control_synthetic.xlsx", v, t, synthetic=True)
    print("built template and synthetic Control workbooks")


if __name__ == "__main__":
    main()
