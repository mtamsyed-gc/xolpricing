"""Stage 06 (stop loss) in the simulator vs an independent numpy implementation (same MINSTD streams, so results are exact).

Runs stages 00-06 (line mode) on a copy of the synthetic Control workbook with 500 iterations (variant T5). The copy
exercises every stop-loss feature: option 1 PCT_BUDGET on gross, option 2 PCT_EPI, option 3 NET_OF_XOL (programme
PP option 1), option 4 AMOUNT with a 20% co-participation; PP option 2 carries a paid reinstatement.
Also checks the statistical sense of the attritional model (mean and CoV of the draws).
"""
from __future__ import annotations

import math
import pickle
import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools/alteryx"))
sys.path.insert(0, str(ROOT / "tests"))
import xol_common as C  # noqa: E402
import build_stage00, build_stage01, build_stage02, build_stage03, build_stage04a, build_stage04, build_stage05, build_stage06  # noqa: E401,E402
import sim  # noqa: E402
import test_stage04 as T4  # noqa: E402
import test_stage05 as T5V  # noqa: E402

N_ITER = 500
OUT = ROOT / "data/output/t5"
CTL = OUT / "XoL_Control_t5.xlsx"
V = C.Variant("T5", OUT / "wf", "..\\..\\..\\..\\", "data/output/t5/XoL_Control_t5.xlsx",
              "data/synthetic/line", "data/synthetic/claimant", "data/output/t5")
M, A = T4.M, T4.A
results = []


def load(name):
    T4.OUT = OUT
    return T4.load(name)


def close(name, got, exp, tol=1e-9):
    got, exp = np.asarray(got, dtype=float), np.asarray(exp, dtype=float)
    ok = got.shape == exp.shape and np.allclose(got, exp, rtol=tol, atol=1e-6, equal_nan=True)
    results.append((name, ok))
    if not ok:
        print(f"  MISMATCH {name}\n    got {got.ravel()[:6]}\n    exp {exp.ravel()[:6]}")


def ty_of(d: pd.Series, ts: pd.Timestamp) -> np.ndarray:
    md = d.dt.strftime("%m-%d") < ts.strftime("%m-%d")
    return (d.dt.year - md.astype(int)).to_numpy()


def nearest(sorted_vals, n, qm):
    return sorted_vals[max(1, math.ceil(qm * n / 1000)) - 1]


def oracle():
    prm = pd.read_excel(CTL, sheet_name="Params", header=3).set_index("param_id").value
    num = lambda k: float(prm[k]) if not pd.isna(prm.get(k)) else None  # noqa: E731
    q_year, epi, inc = int(prm["quote_year"]), num("epi_quote_year"), num("incidence_adj")
    floor, dist = num("sl_cov_floor"), prm["sl_attritional_dist"]
    ts = pd.Timestamp(prm["treaty_start"])
    T4.OUT, T4.CTL = OUT, CTL
    e4 = T4.oracle()
    iters, u, lives = e4["iters"], e4["u"], e4["lives"]
    thr = num("sl_attritional_threshold")
    t_sl = u if thr is None or thr < u else thr

    # attritional history (line mode: totals from the accumulation units)
    units, adj = load("exp_claimants"), load("bc_adjusted")
    tot = units.groupby("ty").amount_q.sum()
    large = adj[adj.adj > t_sl].groupby("ty").amount_q.sum()
    tr = pd.read_excel(CTL, sheet_name="Trend", header=3).sort_values("year").set_index("year")
    z0 = tr.fillna(0)
    t = (z0.weight_client * z0.client_pct + z0.weight_market * z0.market_pct + z0.weight_inflation * z0.inflation_pct
         + z0.residual_pct)
    idx = ((1 + t) * (1 + z0.one_off_adj)).cumprod() / ((1 + t.iloc[0]) * (1 + z0.one_off_adj.iloc[0]))
    trend_f = idx[q_year] / idx
    cp = pd.read_excel(CTL, sheet_name="Completion", header=3)
    cp = cp[(cp.side == "ATTRITIONAL") & (cp.ap_band_from.fillna(0) == 0)]
    pre = cp.groupby("year").selected_factor.prod()
    w = load("bc_weights").set_index("g_ty").weight
    ex = load("exp_exposure_ty").set_index("ty").exposure_ly_full
    yrs = w.index
    attr_raw = tot.reindex(yrs, fill_value=0) - large.reindex(yrs, fill_value=0)
    pppy = attr_raw * trend_f.reindex(yrs) * pre.reindex(yrs, fill_value=1) / ex.reindex(yrs)
    mean_pppy = float((w * pppy).sum())
    cov_hist = pppy.std(ddof=1) / pppy.mean()
    cov = max(cov_hist, floor)
    mu = mean_pppy * lives * inc
    sig = math.sqrt(math.log(1 + cov * cov))
    m_ln = math.log(mu) - sig * sig / 2

    s0 = int(prm["seed"]) % (M - 1) + 1
    st4, st5 = (pow(A, 1 << 28, M) * s0) % M, (pow(A, 3 * (1 << 27), M) * s0) % M
    u4, u5 = T4.minstd(st4, iters), T4.minstd(st5, iters)
    z = np.sqrt(-2 * np.log(u4)) * np.cos(6.283185307179586 * u5)
    attr = np.maximum(mu * (1 + cov * z), 0) if dist == "NORMAL" else np.exp(m_ln + sig * z)
    claim = e4["claim"]
    lgs = np.bincount(e4["it_idx"], weights=np.where(claim > t_sl, claim, 0), minlength=iters)
    gross = attr + lgs
    budget = gross.mean() if num("net_claims_budget") is None else num("net_claims_budget")

    lay = pd.read_excel(CTL, sheet_name="Layers", header=3)
    pg = pd.read_excel(CTL, sheet_name="Programme", header=3)
    pg = pg[(pg.active == 1) & pg.option_id.notna()].sort_values("step")
    first = lambda cv: int(pg[pg.cover == cv].option_id.iloc[0]) if (pg.cover == cv).any() else None  # noqa: E731
    net_pp, net_ev, prog_sl = first("PP"), first("EV"), first("SL")
    e5 = T5V.oracle(OUT, CTL, iters)
    zero = np.zeros(iters)
    ev = {}
    for (o, lyr), d in e5["sim"].items():
        x = ev.setdefault(o, dict(rec=zero.copy(), recn=zero.copy(), ru1=zero.copy(), ru2=zero.copy(), rn1=zero.copy(),
                                  rn2=zero.copy()))
        x["rec"] += d["rec"]
        x["recn"] += d["recn"]
        x[f"ru{lyr}"] += d["ru"]
        x[f"rn{lyr}"] += d["run"]
    structs = {"0 NONE": dict(rec=zero, net=gross, ru=zero, ru_ev=zero)}
    for o, d in e4["out"].items():
        structs[f"1 PP {o}"] = dict(rec=d["rec"], net=gross - d["rec"], ru=d["ru"], ru_ev=zero)
    for o, d in ev.items():
        structs[f"3 EV {o}"] = dict(rec=d["rec"], net=gross - d["rec"], ru=zero, ru_ev=d["ru1"])
    pp_p = e4["out"][net_pp]["rec"] if net_pp else zero
    ev_p = (ev[net_ev]["recn"] if net_pp else ev[net_ev]["rec"]) if net_ev in ev else zero
    sl_rec = {}
    sl_terms = {}
    for _, L in lay[(lay.cover == "SL") & (lay.active == 1) & (lay.layer == 1)].iterrows():
        meas = L.sl_measure if not pd.isna(L.sl_measure) else "PCT_BUDGET"
        base_amt = {"PCT_BUDGET": budget, "PCT_EPI": epi, "AMOUNT": 1.0}[meas]
        ap_amt = L.sl_ap * base_amt
        lim_amt = None if pd.isna(L.sl_limit) else L.sl_limit * base_amt
        netx = L.sl_claims_basis == "NET_OF_XOL"
        base = gross - (pp_p + ev_p if netx else 0)
        over = np.maximum(base - ap_amt, 0)
        cop = 0 if pd.isna(L.co_participation_pct) else L.co_participation_pct
        rec = (over if lim_amt is None else np.minimum(over, lim_amt)) * (1 - cop)
        ru_ev = (ev[net_ev]["rn1"] if net_pp else ev[net_ev]["ru1"]) if netx and net_ev in ev else zero
        structs[f"2 SL {int(L.option_id)}"] = dict(rec=rec, net=base - rec,
                                                   ru=e4["out"][net_pp]["ru"] if netx and net_pp else zero, ru_ev=ru_ev)
        sl_terms[int(L.option_id)] = (ap_amt, lim_amt)
        sl_rec[int(L.option_id)] = rec
    if prm["evaluate_mode"] == "PROGRAMME":
        tot = pp_p + ev_p + (sl_rec[prog_sl] if prog_sl in sl_rec else zero)
        structs["4 PROGRAMME"] = dict(rec=tot, net=gross - tot, ru=e4["out"][net_pp]["ru"] if net_pp else zero,
                                      ru_ev=(ev[net_ev]["rn1"] if net_pp else ev[net_ev]["ru1"]) if net_ev in ev else zero)

    # exposed lives
    mem = load("int_members")
    mem = mem[ty_of(pd.to_datetime(mem.cohort_month), ts) == w.index.max()]
    share = {}
    for _, L in lay[(lay.cover == "PP") & (lay.active == 1)].iterrows():
        ok = mem.annual_limit.isna() | (mem.annual_limit > L.ap)
        share[int(L.option_id)] = mem.cov_days[ok].sum() / mem.cov_days.sum()
    return dict(pppy=pppy, mean_pppy=mean_pppy, cov_hist=cov_hist, cov=cov, mu=mu, attr=attr, lgs=lgs, gross=gross,
                budget=budget, structs=structs, sl_terms=sl_terms, share=share, lives=lives, t_sl=t_sl)


def claimant_mode():
    """Claimant listing + monthly aggregate (no member detail): totals come from the aggregate, all lives exposed."""
    out = ROOT / "data/output/t5c"
    out.mkdir(parents=True, exist_ok=True)
    shutil.copy(CTL, out / "XoL_Control_t5c.xlsx")
    vc = C.Variant("T5C", out / "wf", "..\\..\\..\\..\\", "data/output/t5c/XoL_Control_t5c.xlsx",
                   "data/synthetic/line", "data/synthetic/claimant", "data/output/t5c")
    for b in (build_stage00.build(vc), build_stage01.build_claimant(vc), build_stage02.build(vc),
              build_stage03.build(vc), build_stage04a.build(vc), build_stage04.build(vc), build_stage05.build(vc), build_stage06.build(vc)):
        sim.run(b)
        results.append((f"claimant mode {b.stem}: no STOP", not [m for m in sim.MESSAGES if m[0] == "Error"]))
    with open(out / "stage/int_members.yxdb", "rb") as f:
        results.append(("claimant mode writes an empty int_members", len(pickle.load(f)[1]) == 0))
    with open(out / "stage/int_claims_aggregate.yxdb", "rb") as f:
        fields, rows = pickle.load(f)
    ag = pd.DataFrame(rows, columns=list(fields))
    ag["ty"] = pd.to_datetime(ag.cohort_month).dt.year            # RA, treaty start 1 January
    tot = (ag.amount_paid + ag.amount_os).groupby(ag.ty).sum()
    ay = pd.read_csv(out / "results/out_attritional_by_year.csv").set_index("ty")
    close("claimant mode: total claims from the aggregate", ay.total_claims, tot.reindex(ay.index))
    ck = pd.read_csv(out / "checks/checks_06.csv")
    results.append(("claimant mode: G07 (no member limits, all lives exposed)", "G07" in set(ck.code)))
    el = pd.read_csv(out / "results/out_exposed_lives.csv")
    results.append(("claimant mode: exposed share 1", bool(np.allclose(el.share_exposed, 1))))


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    shutil.copy(ROOT / C.SYNTH.control, CTL)
    wb = load_workbook(CTL)
    ws = wb["Params"]
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value == "iterations":
            ws.cell(r, 3).value = str(N_ITER)
    ly = wb["Layers"]
    hdr = [c.value for c in ly[4]]
    col = lambda n: hdr.index(n) + 1  # noqa: E731
    for r in range(5, ly.max_row + 1):
        cov_, opt = ly.cell(r, 1).value, ly.cell(r, 2).value
        if cov_ == "PP" and opt == 2:
            ly.cell(r, col("reinstatements_n")).value = 1
            ly.cell(r, col("reinstatement_cost_pct")).value = 1.0
        if cov_ == "SL" and opt == 2:
            ly.cell(r, col("sl_measure")).value = "PCT_EPI"
            ly.cell(r, col("sl_ap")).value = 0.95
            ly.cell(r, col("sl_limit")).value = 0.10
        if cov_ == "SL" and opt == 3:
            ly.cell(r, col("sl_claims_basis")).value = "NET_OF_XOL"
            ly.cell(r, col("sl_ap")).value = 0.9
        if cov_ == "SL" and opt == 4:
            ly.cell(r, col("sl_measure")).value = "AMOUNT"
            ly.cell(r, col("sl_ap")).value = 60_000_000
            ly.cell(r, col("sl_limit")).value = 10_000_000
            ly.cell(r, col("co_participation_pct")).value = 0.2
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value == "evaluate_mode":
            ws.cell(r, 3).value = "PROGRAMME"
    wb.save(CTL)
    for b in (build_stage00.build(V), build_stage01.build_line(V), build_stage02.build(V), build_stage03.build(V),
              build_stage04a.build(V), build_stage04.build(V), build_stage05.build(V), build_stage06.build(V)):
        sim.run(b)
        results.append((f"{b.stem}: no STOP", not [m for m in sim.MESSAGES if m[0] == "Error"]))

    compare(oracle(), OUT / "results", OUT / "checks/checks_06.csv")
    # test-copy specifics
    ss = pd.read_csv(OUT / "results/out_struct_summary.csv").set_index("struct_key")
    ck = pd.read_csv(OUT / "checks/checks_06.csv")
    results.append(("NET_OF_XOL option 3 nets PP option 1 and EV option 1",
                    ss.loc["2 SL 3", "pp_option"] == 1 and ss.loc["2 SL 3", "ev_option"] == 1))
    results.append(("PROGRAMME structure: PP 1 + EV 1 + SL 1", "4 PROGRAMME" in ss.index and
                    tuple(ss.loc["4 PROGRAMME", ["pp_option", "ev_option", "sl_option"]]) == (1, 1, 1)))
    results.append(("G08 flags the reinstated PP option", "G08" in set(ck.code)))
    claimant_mode()
    for n, ok in results:
        print(f"{'PASS' if ok else 'FAIL'}  {n}")
    return 0 if all(ok for _, ok in results) else 1


def compare(e, R, checks, tol=1e-9):
    """Outputs in folder R (simulator or Designer) vs the oracle e."""
    rd = lambda f: pd.read_csv(R / f, encoding="utf-8-sig")  # noqa: E731
    ay = rd("out_attritional_by_year.csv").set_index("ty")
    gs = rd("out_gross_summary.csv").iloc[0]
    gi = rd("out_gross_iterations.csv").sort_values("iter")
    ss = rd("out_struct_summary.csv").set_index("struct_key")
    sq = rd("out_struct_quantiles.csv")
    el = rd("out_exposed_lives.csv")
    ck = pd.read_csv(checks, encoding="utf-8-sig")
    n_it = len(gi)
    _close = globals()["close"]

    def close(name, got, exp):  # noqa: F811
        _close(name, got, exp, tol)
    close("attritional cost per life by year", ay.attritional_pppy, e["pppy"].values)
    close("weighted cost per life, CoV", [gs.attritional_pppy_weighted, gs.cov_hist, gs.cov_used],
          [e["mean_pppy"], e["cov_hist"], e["cov"]])
    close("attritional mean for the quote year", [gs.attritional_mean_q], [e["mu"]])
    close("attritional by iteration", gi.attritional, e["attr"])
    close("large claims above t_sl by iteration", gi.large_gross_sl, e["lgs"])
    close("gross by iteration", gi.gross, e["gross"])
    close("budget", [gs.budget], [e["budget"]])
    for k, d in e["structs"].items():
        close(f"{k}: mean / sd recovery, mean / sd net", [ss.loc[k, "mean_recovery"], ss.loc[k, "sd_recovery"],
                                                         ss.loc[k, "mean_net"], ss.loc[k, "sd_net"]],
              [d["rec"].mean(), d["rec"].std(), d["net"].mean(), d["net"].std()])
        g = sq[sq.struct_key == k].sort_values("q")
        qm = np.rint(g.q * 1000).astype(int)
        order = np.lexsort((np.arange(n_it), d["net"]))
        rk = [max(1, math.ceil(x * n_it / 1000)) - 1 for x in qm]
        close(f"{k}: net claims quantiles", g.net_claims, d["net"][order][rk])
        close(f"{k}: reinstatement usage at net quantile", g.ru_l1_at_q, d["ru"][order][rk])
        close(f"{k}: EV reinstatement usage at net quantile", g.ru_ev1_at_q, d["ru_ev"][order][rk])
        close(f"{k}: recovery quantiles", g.recovery, np.sort(d["rec"])[rk])
    for o, (ap_amt, lim_amt) in e["sl_terms"].items():
        k = f"2 SL {o}"
        close(f"SL {o} attachment / limit", [ss.loc[k, "ap_amount"], ss.loc[k, "limit_amount"]],
              [ap_amt, np.nan if lim_amt is None else lim_amt])
    pp = el[el.cover == "PP"].set_index("option_id")
    close("PP exposed share", pp.share_exposed.sort_index(), [e["share"][o] for o in sorted(e["share"])])
    close("exposed lives", pp.exposed_lives.sort_index(), [e["share"][o] * e["lives"] for o in sorted(e["share"])])
    results.append(("SL / EV exposed lives = all lives", bool(np.allclose(el[el.cover != "PP"].share_exposed, 1))))
    # statistical sense and checks
    a = e["attr"]
    results.append((f"attritional draws: mean {a.mean() / e['mu']:.3f} x expected, CoV {a.std() / a.mean():.3f} "
                    f"vs {e['cov']:.3f}", abs(a.mean() / e["mu"] - 1) < 0.02 and abs(a.std() / a.mean() / e["cov"] - 1) < 0.15))
    results.append((f"budget = mean gross {gs.budget:,.0f}; without RI mean net = budget",
                    abs(ss.loc["0 NONE", "mean_net"] - gs.budget) < 1e-6 * gs.budget))
    results.append(("PP limits are known: no G07", "G07" not in set(ck.code)))
    results.append(("no STOP in checks_06", "STOP" not in set(ck.severity)))


if __name__ == "__main__":
    sys.exit(main())
