"""Stage 05 (per-event XoL) in the simulator vs an independent pandas / numpy implementation.

Runs stages 00-03 and 05 (line mode) on a copy of the synthetic Control with 500 iterations (variant T5E): the
synthetic book carries two supplied events (fire 4 lives, road accident 7 lives) and one unlabelled food-poisoning
cluster (18 lives, same scheme and site). The copy adds a second EV option with an AAD and limited reinstatements, and a
second scenario. Events, deterministic burning cost, scenarios and every simulated iteration must equal the oracle
(same MINSTD streams). A SUPPLIED-only run must find the two supplied events only.
"""
from __future__ import annotations

import math
import pickle
import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools/alteryx"))
sys.path.insert(0, str(ROOT / "tests"))
import xol_common as C  # noqa: E402
import build_stage00, build_stage01, build_stage02, build_stage03, build_stage05  # noqa: E401,E402
import sim  # noqa: E402
import test_stage04 as T4  # noqa: E402

N_ITER = 500
OUT = ROOT / "data/output/t5e"
CTL = OUT / "XoL_Control_t5e.xlsx"
M, A, BIG = T4.M, T4.A, 1e15
results = []


def variant(out, ctl):
    return C.Variant(out.name.upper(), out / "wf", "..\\..\\..\\..\\", str(ctl.relative_to(ROOT)).replace("\\", "/"),
                     "data/synthetic/line", "data/synthetic/claimant", str(out.relative_to(ROOT)))


def load(out, name):
    with open(out / "stage" / f"{name}.yxdb", "rb") as f:
        fields, rows = pickle.load(f)
    return pd.DataFrame(rows, columns=list(fields))


def close(name, got, exp, tol=1e-9):
    got, exp = np.asarray(got, dtype=float), np.asarray(exp, dtype=float)
    ok = got.shape == exp.shape and np.allclose(got, exp, rtol=tol, atol=1e-6, equal_nan=True)
    results.append((name, ok))
    if not ok:
        print(f"  MISMATCH {name}\n    got {got.ravel()[:8]}\n    exp {exp.ravel()[:8]}")


def sheet(ctl, name):
    return pd.read_excel(ctl, sheet_name=name, header=3).dropna(how="all")


def ty_of(d: pd.Series, ts: pd.Timestamp) -> np.ndarray:
    return (d.dt.year - (d.dt.strftime("%m-%d") < ts.strftime("%m-%d")).astype(int)).to_numpy()


def layer_loss(loss, ap, lim, cop):
    over = loss - ap
    return np.where(over <= 0, 0, np.minimum(over, lim)) * (1 - cop)


def oracle(out, ctl, n_iter):
    prm = sheet(ctl, "Params").set_index("param_id").value
    num = lambda k: None if pd.isna(prm.get(k)) else float(prm[k])  # noqa: E731
    ts = pd.Timestamp(prm["treaty_start"])
    years = [int(y) for y in str(prm["years_included"]).split(",")]
    qs, inc, full_n = num("qs_cedant_retention_pct"), num("incidence_adj"), num("full_cred_n")
    pol_lim = num("policy_limit")
    src = prm["ev_freq_source"]
    rule = sheet(ctl, "EventRule").iloc[0]
    q = int(prm["quote_year"])

    tr = sheet(ctl, "Trend").sort_values("year").set_index("year").fillna(0)
    t = (tr.weight_client * tr.client_pct + tr.weight_market * tr.market_pct + tr.weight_inflation * tr.inflation_pct
         + tr.residual_pct)
    t = t.where(sheet(ctl, "Trend").sort_values("year").set_index("year").large_claim_trend_pct.isna(), tr.large_claim_trend_pct)
    idx = ((1 + t) * (1 + tr.one_off_adj)).cumprod() / ((1 + t.iloc[0]) * (1 + tr.one_off_adj.iloc[0]))
    trend_f = idx[q] / idx

    c = load(out, "int_claims")
    for d in ("treatment_date", "admission_date"):
        c[d] = pd.to_datetime(c[d])
    c["amt"] = (c.amount_paid_q.fillna(0) + c.amount_os_q.fillna(0)) * qs
    c["ev_date"] = c.admission_date.fillna(c.treatment_date)
    c["ev_day"] = (c.ev_date - pd.Timestamp("2000-01-01")).dt.days
    has_id = c.event_id.notna() & (c.event_id.fillna("").str.strip() != "")
    sup = c[has_id].copy()
    sup["ev_key"] = "ID " + sup.event_id.str.strip().str.upper()
    cand = c[~has_id & c.cause_code.notna() & (c.cause_code.fillna("").str.strip() != "")].copy()
    if rule.id_source == "CLUSTER" and len(cand):
        f = lambda s: s.fillna("?")  # noqa: E731
        gk = {"SCHEME": f(cand.scheme_id), "SITE": f(cand.site), "SCHEME+SITE": f(cand.scheme_id) + "|" + f(cand.site),
              "SCHEME+CAUSE": f(cand.scheme_id)}.get(rule.group_key, pd.Series("*", index=cand.index))
        cand["ckey"] = gk + "|" + cand.cause_code.str.strip().str.upper()
        cand = cand.sort_values(["ckey", "ev_day", "claim_key"], kind="stable")
        start, prev_key, prev = [], None, 0
        for k, d in zip(cand.ckey, cand.ev_day):
            if k != prev_key:
                prev = 0
            prev = d if prev == 0 or d - prev > rule.window_days else prev
            start.append(prev)
            prev_key = k
        cand["ev_key"] = "CL " + cand.ckey + "|" + pd.Series(start, index=cand.index).astype(str)
        el = pd.concat([sup, cand])
    else:
        el = sup
    life = el.groupby(["ev_key", "member_id"]).agg(life_amt=("amt", "sum"), annual_limit=("annual_limit", "max"))
    evs = el.groupby("ev_key").agg(event_date=("ev_date", "min"), n_lives=("member_id", "nunique"))
    evs["ty"] = ty_of(evs.event_date, ts)
    evs = evs[evs.ty.isin(years) & (evs.n_lives >= rule.min_lives)].sort_index()
    life = life.reset_index()
    life = life[life.ev_key.isin(evs.index)].sort_values(["ev_key", "member_id"]).reset_index(drop=True)
    life["ty"] = life.ev_key.map(evs.ty)
    cap = life.annual_limit.fillna(BIG) if pol_lim is None else np.minimum(life.annual_limit.fillna(BIG), pol_lim)
    life["cap"] = cap
    life["adj"] = np.minimum(life.life_amt * life.ty.map(trend_f), cap)

    lay = sheet(ctl, "Layers")
    pg = sheet(ctl, "Programme")
    pg = pg[(pg.cover == "PP") & (pg.active == 1) & pg.option_id.notna()].sort_values("step")
    ppl = lay[(lay.cover == "PP") & (lay.active == 1) & (lay.option_id == pg.option_id.iloc[0])] if len(pg) else lay[:0]

    def pp_rec(cost, capv):
        r = np.zeros_like(cost, dtype=float)
        for _, L in ppl.iterrows():
            if L.limit_basis == "TO_BENEFIT_LIMIT":
                lim = np.where(capv >= BIG, BIG, capv - L.ap)
            else:
                lim = np.full_like(cost, BIG if pd.isna(L.limit) else L.limit, dtype=float)
            cop = 0 if pd.isna(L.co_participation_pct) else L.co_participation_pct
            r += np.where(lim > 0, layer_loss(cost, L.ap, lim, cop), 0)
        return r
    life["net"] = life.adj - pp_rec(life.adj.to_numpy(), life.cap.to_numpy())
    evs["loss_adj"] = life.groupby("ev_key").adj.sum()
    evs["loss_net"] = life.groupby("ev_key").net.sum()

    lives_q = load(out, "bc_lives").lives_quote_year.iloc[0]
    w = load(out, "bc_weights").set_index("g_ty").weight
    ex = load(out, "exp_exposure_ty").set_index("ty").exposure_ly_full.reindex(w.index)
    evl = lay[(lay.cover == "EV") & (lay.active == 1)]
    det = {}
    for _, L in evl.iterrows():
        lim = BIG if pd.isna(L.limit) else L.limit
        cop = 0 if pd.isna(L.co_participation_pct) else L.co_participation_pct
        xl = layer_loss(evs.loss_adj.to_numpy(), L.ap, lim, cop)
        by = pd.Series(xl, index=evs.ty.to_numpy()).groupby(level=0).sum().reindex(w.index, fill_value=0)
        aad = 0 if pd.isna(L.aad) else L.aad
        agg = L.aal if not pd.isna(L.aal) else (BIG if pd.isna(L.reinstatements_n) or pd.isna(L.limit)
                                                 else (1 + L.reinstatements_n) * L.limit)
        rec = np.minimum(np.maximum(by - aad, 0), agg)
        bc = (w * rec / ex).sum()
        n_xl = int((xl > 0).sum())
        det[(int(L.option_id), int(L.layer))] = dict(by=rec, det=bc * lives_q * inc, z=min(1, math.sqrt(n_xl / full_n)))
    n_ty = evs.groupby("ty").size().reindex(w.index, fill_value=0)
    n_hist, n_pool = len(evs), len(life)
    lam = (w * n_ty / ex).sum() * lives_q * inc if src in ("EXPERIENCE", "BOTH") and n_hist > 0 else 0.0

    s0 = int(prm["seed"]) % (M - 1) + 1
    st = {j: (pow(A, o, M) * s0) % M for j, o in {6: 1 << 24, 7: 1 << 25, 8: 1 << 26, 9: 1 << 27}.items()}

    def counts(lmb, u):
        k_max = math.ceil(lmb + 10 * math.sqrt(lmb) + 20)
        pk = [math.exp(-lmb)]
        for k in range(1, k_max + 1):
            pk.append(pk[-1] * lmb / k)
        return np.searchsorted(np.cumsum(pk), u, side="left")
    n_exp = counts(lam, T4.minstd(st[6], n_iter))
    sc = sheet(ctl, "EvScenarios").reset_index(drop=True)
    sc["net_cost"] = sc.cost_per_life - pp_rec(sc.cost_per_life.to_numpy(dtype=float), np.full(len(sc), BIG))
    sim_sc = sc[sc.return_period_yrs.fillna(0) > 0] if src in ("SCENARIOS", "BOTH") else sc[:0]
    u9 = T4.minstd(st[9], n_iter * len(sim_sc)).reshape(n_iter, len(sim_sc)) if len(sim_sc) else np.zeros((n_iter, 0))
    n_sc = np.column_stack([counts(1 / rp, u9[:, k]) for k, rp in enumerate(sim_sc.return_period_yrs)]) \
        if len(sim_sc) else np.zeros((n_iter, 0), dtype=int)

    tot = int(n_exp.sum())
    u7 = T4.minstd(st[7], tot)
    ev_iter = np.repeat(np.arange(n_iter), n_exp)
    k_lives = evs.n_lives.to_numpy()[np.floor(u7 * n_hist).astype(int)] if tot else np.zeros(0, dtype=int)
    u8 = T4.minstd(st[8], int(k_lives.sum()))
    pick = np.floor(u8 * n_pool).astype(int)
    ev_of_life = np.repeat(np.arange(tot), k_lives)
    loss = np.bincount(ev_of_life, weights=life.adj.to_numpy()[pick], minlength=tot)
    loss_net = np.bincount(ev_of_life, weights=life.net.to_numpy()[pick], minlength=tot)
    s_iter = [np.repeat(np.arange(n_iter), n_sc[:, k]) for k in range(len(sim_sc))]
    all_iter = np.concatenate([ev_iter] + s_iter)
    all_loss = np.concatenate([loss] + [np.full(len(si), sim_sc.lives.iloc[k] * sim_sc.cost_per_life.iloc[k])
                                        for k, si in enumerate(s_iter)])
    all_net = np.concatenate([loss_net] + [np.full(len(si), sim_sc.lives.iloc[k] * sim_sc.net_cost.iloc[k])
                                           for k, si in enumerate(s_iter)])
    simres = {}
    for _, L in evl.iterrows():
        lim = BIG if pd.isna(L.limit) else L.limit
        cop = 0 if pd.isna(L.co_participation_pct) else L.co_participation_pct
        xl = np.bincount(all_iter, weights=layer_loss(all_loss, L.ap, lim, cop), minlength=n_iter)
        xln = np.bincount(all_iter, weights=layer_loss(all_net, L.ap, lim, cop), minlength=n_iter)
        aad = 0 if pd.isna(L.aad) else L.aad
        n_r = L.reinstatements_n
        agg = L.aal if not pd.isna(L.aal) else (BIG if pd.isna(n_r) or pd.isna(L.limit) else (1 + n_r) * L.limit)
        rec = np.minimum(np.maximum(xl - aad, 0), agg)
        recn = np.minimum(np.maximum(xln - aad, 0), agg)
        cost = L.reinstatement_cost_pct
        ru = (np.zeros(n_iter) if pd.isna(cost) or pd.isna(L.limit) else
              cost * (rec if pd.isna(n_r) else np.minimum(rec, n_r * L.limit)) / L.limit)
        run = (np.zeros(n_iter) if pd.isna(cost) or pd.isna(L.limit) else
               cost * (recn if pd.isna(n_r) else np.minimum(recn, n_r * L.limit)) / L.limit)
        simres[(int(L.option_id), int(L.layer))] = dict(rec=rec, recn=recn, ru=ru, run=run)
    return dict(evs=evs, life=life, det=det, lam=lam, n_exp=n_exp, n_sc=n_sc, sim=simres, sc=sc, n_hist=n_hist,
                evl=evl, pp_rec=pp_rec)


def prepare(out, ctl, edit):
    out.mkdir(parents=True, exist_ok=True)
    shutil.copy(ROOT / C.SYNTH.control, ctl)
    wb = load_workbook(ctl)
    ws = wb["Params"]
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value == "iterations":
            ws.cell(r, 3).value = str(N_ITER)
    edit(wb)
    wb.save(ctl)
    v = variant(out, ctl)
    for b in (build_stage00.build(v), build_stage01.build_line(v), build_stage02.build(v), build_stage03.build(v),
              build_stage05.build(v)):
        sim.run(b)
        results.append((f"{out.name} {b.stem}: no STOP", not [m for m in sim.MESSAGES if m[0] == "Error"]))


def edit_main(wb):
    ly = wb["Layers"]
    hdr = [c.value for c in ly[4]]
    row = {"cover": "EV", "option_id": 2, "layer": 1, "label": "500k xs 400k, AAD 100k", "limit_basis": "AMOUNT",
           "ap": 400_000, "limit": 500_000, "aad": 100_000, "reinstatements_n": 2, "reinstatement_cost_pct": 0.5,
           "active": 1}
    r = ly.max_row + 1
    for k, val in row.items():
        ly.cell(r, hdr.index(k) + 1).value = val
    sc = wb["EvScenarios"]
    r = sc.max_row + 1
    for c_, val in enumerate(["Hotel fire", 25, 60_000, 100, "DIFC"], 1):
        sc.cell(r, c_).value = val


def compare(e, out, R=None):
    R = R or out / "results"
    ev = pd.read_csv(R / "out_ev_events.csv").set_index("event_key").sort_index()
    results.append((f"events found: {list(ev.index)}", list(ev.index) == list(e["evs"].index)))
    close("event lives", ev.n_lives, e["evs"].n_lives)
    close("event loss (adjusted)", ev.loss_adj, e["evs"].loss_adj)
    close("event loss net of the programme PP option", ev.loss_net_pp, e["evs"].loss_net)
    bc = pd.read_csv(R / "out_ev_bc_summary.csv").set_index(["option_id", "layer"])
    by = pd.read_csv(R / "out_ev_bc_by_year.csv")
    sm = pd.read_csv(R / "out_ev_sim_summary.csv").set_index(["option_id", "layer"])
    its = pd.read_csv(R / "out_ev_iterations.csv")
    for k, d in e["det"].items():
        close(f"EV {k} recovery by TY", by[(by.option_id == k[0]) & (by.layer == k[1])].sort_values("ty").recovery, d["by"])
        close(f"EV {k} expected cost (det.), Z", [bc.loc[k, "expected_cost_det"], bc.loc[k, "credibility_z"]],
              [d["det"], d["z"]])
        g = its[(its.option_id == k[0]) & (its.layer == k[1])].sort_values("iter")
        s = e["sim"][k]
        close(f"EV {k} recovery by iteration", g.recovery, s["rec"])
        close(f"EV {k} recovery net of PP by iteration", g.recovery_net, s["recn"])
        close(f"EV {k} reinstatement usage by iteration", g.reinst_usage, s["ru"])
        close(f"EV {k} reinstatement usage (net of PP) by iteration", g.reinst_usage_net, s["run"])
        close(f"EV {k} mean / sd", [sm.loc[k, "mean_recovery"], sm.loc[k, "sd_recovery"]],
              [s["rec"].mean(), s["rec"].std()])
        z = d["z"]
        close(f"EV {k} blended", [sm.loc[k, "expected_loss_blended"]], [z * d["det"] + (1 - z) * s["rec"].mean()])
    close("experience lambda", [bc.lambda_exp.iloc[0]], [e["lam"]])
    n_tot = e["n_exp"] + e["n_sc"].sum(axis=1)
    g = its[its.option_id == its.option_id.min()].sort_values("iter")
    close("events per iteration", g.n_events, n_tot)
    sc = pd.read_csv(R / "out_ev_scenarios.csv")
    for _, L in e["evl"].iterrows():
        lim = BIG if pd.isna(L.limit) else L.limit
        cop = 0 if pd.isna(L.co_participation_pct) else L.co_participation_pct
        s_ = sc[sc.option_id == L.option_id].set_index("scenario_name").reindex(e["sc"].scenario_name)
        close(f"scenario layer loss EV {int(L.option_id)}", s_.layer_loss,
              layer_loss(e["sc"].lives * e["sc"].cost_per_life, L.ap, lim, cop))


def main() -> int:
    prepare(OUT, CTL, edit_main)
    e = oracle(OUT, CTL, N_ITER)
    compare(e, OUT)
    n = e["n_exp"]
    results.append((f"experience counts mean {n.mean():.3f} vs lambda {e['lam']:.3f}", abs(n.mean() / e["lam"] - 1) < 0.15))
    # SUPPLIED only: the unlabelled food-poisoning cluster is not an event
    out2 = ROOT / "data/output/t5s"
    ctl2 = out2 / "XoL_Control_t5s.xlsx"

    def edit_sup(wb):
        edit_main(wb)
        wb["EventRule"].cell(5, 1).value = "SUPPLIED"
    prepare(out2, ctl2, edit_sup)
    e2 = oracle(out2, ctl2, N_ITER)
    compare(e2, out2)
    results.append(("SUPPLIED: 2 events", e2["n_hist"] == 2))
    for nm, ok in results:
        print(f"{'PASS' if ok else 'FAIL'}  {nm}")
    return 0 if all(ok for _, ok in results) else 1


if __name__ == "__main__":
    sys.exit(main())
