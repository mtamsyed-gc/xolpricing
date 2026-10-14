"""The simulator must reproduce the stored Designer outputs (tests/reconciliation/designer_twin)."""
import sys
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools/alteryx"))
import xol_common as C  # noqa: E402
import build_stage00  # noqa: E402
import build_stage01  # noqa: E402
import build_stage02  # noqa: E402
import build_stage03  # noqa: E402
import numpy as np  # noqa: E402
import sim  # noqa: E402

TWIN = ROOT / "tests/reconciliation/designer_twin"
OUT = ROOT / "data/output/synth/checks"


def main() -> int:
    ok = True
    sim.run(build_stage00.build(C.SYNTH))
    sim.run(build_stage01.build_line(C.SYNTH))
    for f in ("checks_00", "out_scenario_applied", "checks_01a", "recon_01a"):
        ok &= compare(f)
    sim.run(build_stage02.build(C.SYNTH))
    sim.run(build_stage03.build(C.SYNTH))
    for f in ("checks_02", "recon_02", "out_completion_indicated", "checks_03", "recon_03"):
        ok &= compare(f)
    for f in ("out_bc_summary", "out_bc_by_year", "out_annualisation"):
        ok &= compare_table(f)
    sim.run(build_stage01.build_claimant(C.SYNTH))
    for f in ("checks_01b", "recon_01b"):
        ok &= compare(f)
    sim.run(build_stage01.build_ete(C.SYNTH))
    for f in ("checks_01c", "recon_01c"):
        ok &= compare(f)
    ok &= stage04()
    ok &= twin("stage05", "twin_stage05.py", "stage 05 (per-event)")
    ok &= stage06()
    ok &= stage07()
    print("PASS" if ok else "FAIL")
    return 0 if ok else 1


def stage04() -> bool:
    """Designer's 10,000-iteration stage-04 run vs the numpy oracle (exact RNG replica)."""
    import subprocess
    import gzip
    import shutil
    d = TWIN / "stage04"
    tmp = ROOT / "data/output/twin04"
    tmp.mkdir(parents=True, exist_ok=True)
    for f in ("out_sim_summary.csv", "out_sim_quantiles.csv"):
        shutil.copy(d / f, tmp / f)
    with gzip.open(d / "out_sim_iterations.csv.gz", "rb") as a, open(tmp / "out_sim_iterations.csv", "wb") as b:
        shutil.copyfileobj(a, b)
    r = subprocess.run([sys.executable, str(ROOT / "tests/twin_stage04.py"), str(tmp)], capture_output=True, text=True)
    good = r.returncode == 0
    print(f"{'PASS' if good else 'FAIL'}  stage 04, 10,000 iterations (Designer 2026.2 vs oracle)")
    return good


def stage06() -> bool:
    """Designer's 10,000-iteration stage-06 run (after stage 04 at 10,000) vs the numpy oracle."""
    import subprocess
    import gzip
    import shutil
    d = TWIN / "stage06"
    tmp = ROOT / "data/output/twin06"
    tmp.mkdir(parents=True, exist_ok=True)
    for f in d.iterdir():
        if f.suffix == ".gz":
            with gzip.open(f, "rb") as a, open(tmp / f.stem, "wb") as b:
                shutil.copyfileobj(a, b)
        elif f.suffix == ".csv":
            shutil.copy(f, tmp / f.name)
    r = subprocess.run([sys.executable, str(ROOT / "tests/twin_stage06.py"), str(tmp)], capture_output=True, text=True)
    good = r.returncode == 0
    print(f"{'PASS' if good else 'FAIL'}  stage 06, 10,000 iterations (Designer 2026.2 vs oracle)")
    return good


def unpack(*subdirs) -> Path:
    """Designer's CSVs of the given baseline subfolders (gz expanded) in one folder."""
    import gzip
    import shutil
    tmp = ROOT / "data/output" / ("twin_" + "_".join(subdirs))
    tmp.mkdir(parents=True, exist_ok=True)
    for sd in subdirs:
        for f in (TWIN / sd).iterdir():
            if f.suffix == ".gz":
                with gzip.open(f, "rb") as a, open(tmp / f.stem, "wb") as b:
                    shutil.copyfileobj(a, b)
            elif f.suffix == ".csv":
                shutil.copy(f, tmp / f.name)
    return tmp


def twin(sd, script, label) -> bool:
    import subprocess
    r = subprocess.run([sys.executable, str(ROOT / "tests" / script), str(unpack(sd))], capture_output=True, text=True)
    good = r.returncode == 0
    if not good:
        print(r.stdout[-3000:], r.stderr[-3000:])
    print(f"{'PASS' if good else 'FAIL'}  {label}, 10,000 iterations (Designer 2026.2 vs oracle)")
    return good


def stage07() -> bool:
    """Designer's Results workbook (stage 07 after 04-06 at 10,000 iterations): every data sheet present with the
    rows of Designer's result CSVs; recalculated by LibreOffice, the exhibits equal the pandas computation."""
    import shutil
    from openpyxl import load_workbook
    sys.path.insert(0, str(ROOT / "tests"))
    sys.path.insert(0, str(ROOT / "tools/excel"))
    import results_spec as RS
    import test_stage07 as T7
    d = unpack("stage04", "stage05", "stage06")
    for f in ("out_bc_by_year.csv", "out_bc_summary.csv"):     # stage 03 (Designer), for the as-if exhibit
        shutil.copy(TWIN / f, d / f)
    book = ROOT / "data/output/twin07.xlsx"
    shutil.copy(TWIN / "stage07/XoL_Results_designer.xlsx", book)
    wb = load_workbook(book)      # not read-only: Designer writes no sheet dimension, so max_row needs a full load
    good = set(RS.SHEETS) <= set(wb.sheetnames)
    for sheet in RS.SHEETS:
        if good and (d / f"{sheet}.csv").exists():
            good &= wb[sheet].max_row - 1 == len(pd.read_csv(d / f"{sheet}.csv"))
    rc = T7.recalc(book)
    good &= rc.get("status") == "success" and rc.get("total_errors") == 0
    T7.results.clear()
    T7.exhibit_checks(book, d, ROOT / C.SYNTH.control, tag="Designer: ")
    good &= all(g for _, g in T7.results) and len(T7.results) >= 6
    chk = pd.read_csv(TWIN / "stage07/checks_07.csv")
    good &= not (chk.severity == "STOP").any()
    print(f"{'PASS' if good else 'FAIL'}  stage 07 Results workbook ({len(T7.results)} exhibit checks, "
          f"{rc.get('total_formulas')} formulas, Designer 2026.2 vs pandas)")
    return bool(good)


def compare_table(f) -> bool:
    d, s = pd.read_csv(TWIN / f"{f}.csv"), pd.read_csv(ROOT / "data/output/synth/results" / f"{f}.csv")
    good = set(d.columns) == set(s.columns) and len(d) == len(s)       # columns matched by name
    for c in d.columns if good else []:
        if pd.api.types.is_numeric_dtype(d[c]):
            good &= bool(np.allclose(d[c], s[c], rtol=1e-9, atol=1e-6, equal_nan=True))
        else:
            good &= bool((d[c].fillna("").astype(str) == s[c].fillna("").astype(str)).all())
    print(f"{'PASS' if good else 'FAIL'}  {f} (Designer 2026.2 vs simulator)")
    return good


def compare(f) -> bool:
    d, s = pd.read_csv(TWIN / f"{f}.csv"), pd.read_csv(OUT / f"{f}.csv")
    if f == "out_completion_indicated":
        good = d[["ty", "component", "measure", "months_elapsed"]].equals(s[["ty", "component", "measure", "months_elapsed"]]) \
            and np.allclose(d.indicated_factor, s.indicated_factor, rtol=1e-9)
        print(f"{'PASS' if good else 'FAIL'}  {f} (Designer 2026.2 vs simulator)")
        return bool(good)
    if f.startswith("checks") or f == "out_scenario_applied":
        good = d.equals(s)
    else:
        m = d.merge(s, on=["stage", "table", "measure"], how="outer", suffixes=("_d", "_s"))
        good = m.notna().all().all() and ((m.value_d - m.value_s).abs() <= 1e-4).all()
    print(f"{'PASS' if good else 'FAIL'}  {f} (Designer 2026.2 vs simulator)")
    return bool(good)


if __name__ == "__main__":
    sys.exit(main())
