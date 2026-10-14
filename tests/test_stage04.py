"""Stage 04 in the simulator vs an independent numpy implementation (same MINSTD streams, so results are exact).

Runs stages 00-04 on a copy of the synthetic Control workbook with 500 iterations (variant T4) for speed.
Also checks the statistical sense of the result: alpha near the generator's 1.6, the RNG is uniform, and the
stochastic mean is close to the deterministic burning cost.
"""
from __future__ import annotations

import math
import pickle
import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools/alteryx"))
import xol_common as C  # noqa: E402
import build_stage00, build_stage01, build_stage02, build_stage03, build_stage04  # noqa: E401,E402
import sim  # noqa: E402

N_ITER = 500
OUT = ROOT / "data/output/t4"
CTL = OUT / "XoL_Control_t4.xlsx"
V = C.Variant("T4", OUT / "wf", "..\\..\\..\\..\\", "data/output/t4/XoL_Control_t4.xlsx",
              "data/synthetic/line", "data/synthetic/claimant", "data/output/t4")
M, A = 2147483647, 48271
results = []


def load(name):
    with open(OUT / "stage" / f"{name}.yxdb", "rb") as f:
        fields, rows = pickle.load(f)
    return pd.DataFrame(rows, columns=list(fields))


def close(name, got, exp, tol=1e-9):
    got, exp = np.asarray(got, dtype=float), np.asarray(exp, dtype=float)
    ok = got.shape == exp.shape and np.allclose(got, exp, rtol=tol, atol=1e-6, equal_nan=True)
    results.append((name, ok))
    if not ok:
        print(f"  MISMATCH {name}\n    got {got.ravel()[:6]}\n    exp {exp.ravel()[:6]}")


def minstd(start, n):
    out, x = np.empty(n), start
    for i in range(n):
        x = (A * x) % M
        out[i] = x
    return out / M


def oracle():
    prm = pd.read_excel(CTL, sheet_name="Params", header=3).set_index("param_id").value
    seed, iters = int(prm["seed"]), int(prm["iterations"])
    inc = float(prm["incidence_adj"])
    lay = pd.read_excel(CTL, sheet_name="Layers", header=3)
    lay = lay[(lay.cover == "PP") & (lay.active == 1)].copy()
    u = float(prm["fit_cutoff"]) if not pd.isna(prm.get("fit_cutoff")) else lay.ap.min()
    s0 = seed % (M - 1) + 1
    starts = {j: (pow(A, j * (1 << 29), M) * s0) % M for j in (1, 2, 3)}

    adj = load("bc_adjusted")
    fu = adj[adj.adj > u].copy()
    unc = fu.adj_uncapped < fu.cap
    alpha = unc.sum() / np.log(fu.adj / u).sum()
    caps = fu.sort_values("acc_key").cap.to_numpy()
    big_m = fu.adj.max()

    w = load("bc_weights").set_index("g_ty").weight
    ex = load("exp_exposure_ty").set_index("ty").exposure_ly_full
    n_u = fu.groupby("ty").post_cnt_f_.sum()
    freq = (n_u.reindex(w.index, fill_value=0) / ex.reindex(w.index)) * w
    lives = load("bc_summary").lives_quote_year.max()
    lam = freq.sum() * lives * inc

    n_y, e_y = n_u.reindex(w.index, fill_value=0), ex.reindex(w.index)
    m_pool = n_y.sum() / e_y.sum()
    phi = ((n_y - e_y * m_pool) ** 2 / (e_y * m_pool)).sum() / (len(w) - 1) if len(w) > 1 else 1.0
    nb = prm["freq_model"] == "NEGBIN" and phi > 1
    k_max = math.ceil(lam + 10 * math.sqrt(lam * (phi if nb else 1)) + 20)
    if nb:
        r, qq = lam / (phi - 1), 1 / phi
        pk = [math.exp(r * math.log(qq))]
        for k in range(1, k_max + 1):
            pk.append(pk[-1] * (r + k - 1) / k * (1 - qq))
    else:
        pk = [math.exp(-lam)]
        for k in range(1, k_max + 1):
            pk.append(pk[-1] * lam / k)
    cdf = np.cumsum(pk)
    u1 = minstd(starts[1], iters)
    n_claims = np.searchsorted(cdf, u1, side="left")       # #{k: cdf_k < U}
    tot = int(n_claims.sum())
    u2, u3 = minstd(starts[2], tot), minstd(starts[3], tot)
    x = u * u2 ** (-1 / alpha)
    cap = caps[np.floor(u3 * len(caps)).astype(int)]
    claim = np.minimum(x, cap)
    below = np.minimum(claim, big_m)
    it_idx = np.repeat(np.arange(iters), n_claims)

    out = {}
    for _, L in lay.iterrows():
        if L.limit_basis == "TO_BENEFIT_LIMIT":
            lim = np.where(cap >= 1e15, 1e15, cap - L.ap)
        else:
            lim = np.full_like(claim, L.limit if not pd.isna(L.limit) else 1e15)
        cop = 0 if pd.isna(L.co_participation_pct) else L.co_participation_pct

        def loss(c):
            over = c - L.ap
            return np.where(over <= 0, 0, np.where(over < lim, over, np.where(lim > 0, lim, 0))) * (1 - cop)
        xl = np.bincount(it_idx, weights=loss(claim), minlength=iters)
        xlb = np.bincount(it_idx, weights=loss(below), minlength=iters)
        aad = 0 if pd.isna(L.aad) else L.aad
        n_r = L.reinstatements_n
        agg = L.aal if not pd.isna(L.aal) else (1e15 if pd.isna(n_r) or pd.isna(L.limit) else (1 + n_r) * L.limit)
        rec = np.minimum(np.maximum(xl - aad, 0), agg)
        f = np.where(xl > 0, rec / np.where(xl > 0, xl, 1), 1)
        cost = L.reinstatement_cost_pct
        ru = (np.zeros(iters) if pd.isna(cost) or pd.isna(L.limit) else
              cost * (rec if pd.isna(n_r) else np.minimum(rec, n_r * L.limit)) / L.limit)
        out[int(L.option_id)] = dict(rec=rec, above=(xl - xlb) * f, ru=ru)
    return dict(alpha=alpha, lam=lam, phi=phi, nb=nb, n_claims=n_claims, u1=u1, out=out, claims_sum=claim.sum(), claim=claim,
                it_idx=it_idx, iters=iters, u=u, lives=lives)


def negbin():
    """freq_model NEGBIN: the synthetic yearly counts are overdispersed (phi > 1), so the count draws follow a negative
    binomial with mean lambda and variance phi x lambda; exact match to the oracle."""
    global OUT, CTL
    keep = OUT, CTL
    OUT = ROOT / "data/output/t4n"
    CTL = OUT / "XoL_Control_t4n.xlsx"
    OUT.mkdir(parents=True, exist_ok=True)
    shutil.copy(keep[1], CTL)
    wb = load_workbook(CTL)
    ws = wb["Params"]
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value == "freq_model":
            ws.cell(r, 3).value = "NEGBIN"
    wb.save(CTL)
    vn = C.Variant("T4N", OUT / "wf", "..\\..\\..\\..\\", "data/output/t4n/XoL_Control_t4n.xlsx",
                   "data/synthetic/line", "data/synthetic/claimant", "data/output/t4n")
    for b in (build_stage00.build(vn), build_stage01.build_line(vn), build_stage02.build(vn), build_stage03.build(vn),
              build_stage04.build(vn)):
        sim.run(b)
    e = oracle()
    summ = pd.read_csv(OUT / "results/out_sim_summary.csv").set_index("option_id")
    its = pd.read_csv(OUT / "results/out_sim_iterations.csv")
    results.append((f"NEGBIN used (phi {e['phi']:.2f})", e["nb"] and (summ.freq_dist == "NEGBIN").all()))
    close("NEGBIN total claims", [pd.read_csv(OUT / "checks/recon_04.csv").set_index(["table", "measure"]).value[
        ("iterations", "claims")]], [e["n_claims"].sum()])
    for o, d in e["out"].items():
        g = its[its.option_id == o].sort_values("iter")
        close(f"NEGBIN option {o} recovery by iteration", g.recovery, d["rec"])
    n = e["n_claims"]
    results.append((f"NEGBIN counts: mean {n.mean():.1f} vs lambda {e['lam']:.1f}, variance / mean {n.var() / n.mean():.2f}"
                    f" vs phi {e['phi']:.2f}", abs(n.mean() / e["lam"] - 1) < 0.05 and abs(n.var() / n.mean() / e["phi"] - 1) < 0.3))
    OUT, CTL = keep


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    shutil.copy(ROOT / C.SYNTH.control, CTL)
    wb = load_workbook(CTL)
    ws = wb["Params"]
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value == "iterations":
            ws.cell(r, 3).value = str(N_ITER)
    ly = wb["Layers"]                       # exercise the annual terms: AAD on option 1, 1 paid reinstatement on 2
    hdr = [c.value for c in ly[4]]
    for r in range(5, ly.max_row + 1):
        if ly.cell(r, 1).value == "PP" and ly.cell(r, 2).value == 1:
            ly.cell(r, hdr.index("aad") + 1).value = 1_000_000
        if ly.cell(r, 1).value == "PP" and ly.cell(r, 2).value == 2:
            ly.cell(r, hdr.index("reinstatements_n") + 1).value = 1
            ly.cell(r, hdr.index("reinstatement_cost_pct") + 1).value = 1.0
    wb.save(CTL)
    for b in (build_stage00.build(V), build_stage01.build_line(V), build_stage02.build(V), build_stage03.build(V),
              build_stage04.build(V)):
        sim.run(b)
        results.append((f"{b.stem}: no STOP", not [m for m in sim.MESSAGES if m[0] == "Error"]))

    e = oracle()
    summ = pd.read_csv(OUT / "results/out_sim_summary.csv").set_index("option_id")
    its = pd.read_csv(OUT / "results/out_sim_iterations.csv")
    qs = pd.read_csv(OUT / "results/out_sim_quantiles.csv")
    rec = pd.read_csv(OUT / "checks/recon_04.csv").set_index(["table", "measure"]).value
    close("alpha", [summ.alpha.iloc[0]], [e["alpha"]])
    close("lambda", [summ["lambda"].iloc[0]], [e["lam"]])
    close("total claims", [rec[("iterations", "claims")]], [e["n_claims"].sum()])
    close("total claim amount", [rec[("claims", "amount")]], [e["claims_sum"]])
    for o, d in e["out"].items():
        g = its[its.option_id == o].sort_values("iter")
        close(f"option {o} recovery by iteration", g.recovery, d["rec"])
        close(f"option {o} mean / sd", [summ.loc[o, "mean_recovery"], summ.loc[o, "sd_recovery"]],
              [d["rec"].mean(), d["rec"].std()])
        srt = np.sort(d["rec"])
        q = qs[qs.option_id == o].set_index("q").recovery
        exp_q = [srt[max(1, math.ceil(round(qq, 3) * N_ITER)) - 1] for qq in q.index]
        close(f"option {o} quantiles", q.values, exp_q)
        close(f"option {o} above largest observed", [summ.loc[o, "mean_above_max_observed"]], [d["above"].mean()])
        close(f"option {o} reinstatement usage", [summ.loc[o, "mean_reinst_usage"]], [d["ru"].mean()])
    close("dispersion phi", [summ.phi.iloc[0]], [e["phi"]])
    negbin()
    # statistical sense
    a = summ.alpha.iloc[0]
    results.append((f"alpha {a:.3f} within 1.6 +/- 0.15", abs(a - 1.6) < 0.15))
    u1 = e["u1"]
    results.append((f"iteration uniforms mean {u1.mean():.3f}, KS vs U(0,1) small",
                    abs(u1.mean() - 0.5) < 0.03 and np.max(np.abs(np.sort(u1) - np.arange(1, N_ITER + 1) / N_ITER)) < 0.07))
    results.append((f"reinstatement usage used on option 2 ({summ.loc[2, 'mean_reinst_usage']:.3f})",
                    summ.loc[2, "mean_reinst_usage"] > 0))
    r = summ.det_to_stoch_below[[2, 3, 4]]           # option 1 carries a test AAD
    results.append((f"deterministic / stochastic within 0.8-1.25 ({', '.join(f'{x:.2f}' for x in r)})",
                    bool(((r > 0.8) & (r < 1.25)).all())))
    for n, ok in results:
        print(f"{'PASS' if ok else 'FAIL'}  {n}")
    return 0 if all(ok for _, ok in results) else 1


if __name__ == "__main__":
    sys.exit(main())
