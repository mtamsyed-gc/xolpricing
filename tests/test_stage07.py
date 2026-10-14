"""Stage 07 (Results workbook) end to end: stages 00-07 in the simulator (line mode, 300 iterations, PROGRAMME mode),
then the workbook recalculated by LibreOffice. Checks:
  - every data sheet holds its staged table, with the columns of results_spec (and nothing missing);
  - zero formula errors; Pricing premiums, C1 "1 in N" and C2 surpluses equal an independent pandas computation
    from the result CSVs; a loading override flows through;
  - a second run with scenario_id = 1 keeps scenario 0's rows in scn_summary and adds scenario 1's.
"""
from __future__ import annotations

import json
import math
import pickle
import shutil
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools/alteryx"))
sys.path.insert(0, str(ROOT / "tools/excel"))
import xol_common as C  # noqa: E402
import results_spec as RS  # noqa: E402
import build_results  # noqa: E402
import build_stage00, build_stage01, build_stage02, build_stage03, build_stage04a, build_stage04, build_stage05, build_stage06, build_stage07  # noqa: E401,E402
import sim  # noqa: E402

OUT = ROOT / "data/output/t7"
CTL = OUT / "XoL_Control_t7.xlsx"
V = C.Variant("T7", OUT / "wf", "..\\..\\..\\..\\", "data/output/t7/XoL_Control_t7.xlsx", "data/synthetic/line",
              "data/synthetic/claimant", "data/output/t7")
RECALC = Path("/mnt/skills/public/xlsx/scripts/recalc.py")
results = []


def ok(name, cond):
    results.append((name, bool(cond)))


def close(name, got, exp, tol=1e-7):
    got, exp = np.asarray(got, dtype=float), np.asarray(exp, dtype=float)
    good = got.shape == exp.shape and np.allclose(got, exp, rtol=tol, atol=1e-4, equal_nan=True)
    results.append((name, good))
    if not good:
        print(f"  MISMATCH {name}\n    got {got.ravel()[:8]}\n    exp {exp.ravel()[:8]}")


def set_param(pid, value):
    wb = load_workbook(CTL)
    ws = wb["Params"]
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, 1).value == pid:
            ws.cell(r, 3).value = value
    wb.save(CTL)


def run_all():
    for b in (build_stage00.build(V), build_stage01.build_line(V), build_stage02.build(V), build_stage03.build(V),
              build_stage04a.build(V), build_stage04.build(V), build_stage05.build(V), build_stage06.build(V), build_stage07.build(V)):
        sim.run(b)
        ok(f"{b.stem}: no STOP", not [m for m in sim.MESSAGES if m[0] == "Error"])


def recalc(path):
    r = subprocess.run([sys.executable, str(RECALC), str(path), "180"], capture_output=True, text=True)
    return json.loads(r.stdout)


def staged(name):
    with open(OUT / "stage" / f"{name}.yxdb", "rb") as f:
        fields, rows = pickle.load(f)
    return list(fields), len(rows)


def exhibit_checks(calc, R, ctl, tag=""):
    """Pricing, C1, C2 and Expiring of a recalculated workbook `calc` against an independent pandas computation from
    the result CSVs in R and the Loadings of the Control workbook ctl. Returns premium(cover, option, load)."""
    ss = pd.read_csv(R / "out_struct_summary.csv").set_index("struct_key")
    sq = pd.read_csv(R / "out_struct_quantiles.csv")
    ps = pd.read_csv(R / "out_sim_summary.csv")
    es = pd.read_csv(R / "out_ev_sim_summary.csv")
    lo = pd.read_excel(ctl, sheet_name="Loadings", header=3).set_index("cover")
    gs = pd.read_csv(R / "out_gross_summary.csv").iloc[0]
    el = pd.read_csv(R / "out_exposed_lives.csv")

    def premium(cover, o, load=None):
        L = lo.loc[cover]
        tot = L.reinsurer_expense_pct + L.brokerage_pct + L.profit_margin_pct + L.premium_tax_pct
        if load:
            tot += load
        key = {"PP": "1 PP ", "EV": "3 EV ", "SL": "2 SL "}[cover] + str(o)
        sd = ss.loc[key, "sd_recovery"]
        if cover == "SL":
            e, r, w1 = ss.loc[key, "mean_recovery"], 0.0, 1.0
        else:
            t = ps if cover == "PP" else es
            t = t[t.option_id == o]
            e = t.expected_loss_blended.sum()
            r = (t.mean_recovery * t.mean_reinst_usage).sum() / t.mean_recovery.sum()
            w1 = t[t.layer == 1].mean_recovery.sum() / t.mean_recovery.sum()
        tgt = (e + L.risk_load_k * sd) / (1 - tot)
        p = tgt / (1 + r)
        return dict(e=e, p=p, r=r, p1=p * w1, p2=p * (1 - w1), key=key)

    w = load_workbook(calc, data_only=True)
    pr = w["Pricing"]
    exp_p, got_p = [], []
    opts = [(c, o) for c in ("PP", "EV", "SL") for o in range(1, 5)
            if f"{ {'PP': '1 PP ', 'EV': '3 EV ', 'SL': '2 SL '}[c]}{o}" in ss.index]
    for c, o in opts:
        r = build_results.PROW[(c, o)]
        got_p.append(pr[f"{build_results.PC['Upfront annual premium']}{r}"].value)
        exp_p.append(premium(c, o)["p"])
    close(f"{tag}Pricing upfront premium, {len(opts)} options", got_p, exp_p)
    # 1 in N for PP 1
    p1 = premium("PP", 1)
    q = sq[sq.struct_key == "1 PP 1"]
    n_exp = len(q) / (q.recovery > p1["p"]).sum()
    close(f"{tag}C1 recoveries exceed premium 1 in N (PP 1)",
          [w["Reinsurance Structures"][f"J{5}"].value], [n_exp])
    lives = el[(el.cover == "PP") & (el.option_id == 1)].exposed_lives.iloc[0]
    close(f"{tag}C1 premium per exposed life (PP 1)", [w["Reinsurance Structures"]["F5"].value], [p1["p"] / lives])
    # C2: average year and 1-in-50 for PP 1, EV 1, programme
    ci = w["Cedant Impact"]
    budget = gs.budget
    keys = [ci.cell(5, c).value for c in range(3, 3 + len(build_results.STRUCTS))]
    for key in ("0 NONE", "1 PP 1", "3 EV 1", "2 SL 1", "4 PROGRAMME"):
        if key not in ss.index or key not in keys:
            continue
        col = 3 + keys.index(key)
        s_ = ss.loc[key]
        comp = []
        if s_.pp_option > 0:
            comp.append(("pp", premium("PP", int(s_.pp_option))))
        if s_.ev_option > 0:
            comp.append(("ev", premium("EV", int(s_.ev_option))))
        if s_.sl_option > 0:
            comp.append(("sl", premium("SL", int(s_.sl_option))))
        up = sum(x["p"] for _, x in comp)
        ri_mean = sum(x["p1"] * s_[f"mean_ru_{'l' if k == 'pp' else 'ev'}1"] + x["p2"] * s_[f"mean_ru_{'l' if k == 'pp' else 'ev'}2"]
                      for k, x in comp if k != "sl")
        avg = (budget - s_.mean_net - up - ri_mean) / 1000
        qq = sq[(sq.struct_key == key) & (sq.qm == 980)].iloc[0]
        ri_q = sum(x["p1"] * qq[f"ru_{'l' if k == 'pp' else 'ev'}1_at_q"] + x["p2"] * qq[f"ru_{'l' if k == 'pp' else 'ev'}2_at_q"]
                   for k, x in comp if k != "sl")
        rp50 = (budget - qq.net_claims - up - ri_q) / 1000
        close(f"{tag}C2 {key}: average year, 1 in 50", [ci.cell(11, col).value, ci.cell(13, col).value], [avg, rp50])
    fc = pd.read_csv(R / "out_fit_candidates.csv") if (R / "out_fit_candidates.csv").exists() else pd.DataFrame()
    if "Severity" in w.sheetnames and len(fc):      # (a 0.4 workbook has no Severity sheet)
        sv = w["Severity"]
        got = [[sv.cell(7 + i, c).value for c in (1, 3, 7)] for i in range(len(fc))]
        ok(f"{tag}Severity: candidates, parameter 1 and AIC as out_fit_candidates; used {sv['H4'].value}",
           [g[0] for g in got] == list(fc.dist) and
           np.allclose([g[1] for g in got], fc.p1, rtol=1e-12) and np.allclose([g[2] for g in got], fc.aic, rtol=1e-12)
           and sv["H4"].value == ps.sev_dist.iloc[0] == fc[fc.selected == 1].dist.iloc[0])
    ex = w["Expiring"]
    ok(f"{tag}Expiring: matched option {ex['G5'].value}, actual recoveries read ({ex['H13'].value})",
       ex["G5"].value == 1 and ex["H13"].value == 2650000)
    # as-if: premium at the expiring rate trended one year, on each year's exposure (quote-year level, as the layer loss)
    by = pd.read_csv(R / "out_bc_by_year.csv")
    ctl_ex = pd.read_excel(ctl, sheet_name="Expiring", header=3).dropna(subset=["cover"]).iloc[0]
    bs = pd.read_csv(R / "out_bc_summary.csv").set_index(["option_id", "layer"])
    lev = bs.loc[(1, 1), "lev_trend"]
    q = by[(by.option_id == 1) & (by.layer == 1)].sort_values("ty")
    q_ = int(pd.read_excel(ctl, sheet_name="Params", header=3).set_index("param_id").value["quote_year"])
    got = [[ex.cell(r, c).value for c in (1, 4, 5, 7)] for r in range(13, 13 + len(q))]
    exp = [[t.ty, t.xl_aa, ctl_ex.rate * (1 + lev) * t.exposure_ly_full,
            ctl_ex.rate * (1 + lev) * t.exposure_ly_full / (1 + lev) ** (q_ - t.ty)] for t in q.itertuples()]
    close(f"{tag}Expiring as-if by year: recoveries, premium (quote-year level), premium (year level)", got, exp)

    return premium


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    shutil.copy(ROOT / C.SYNTH.control, CTL)
    set_param("iterations", "300")
    set_param("evaluate_mode", "PROGRAMME")
    build_results.build()
    import zipfile
    rels = zipfile.ZipFile(ROOT / "excel/templates/XoL_Results.xlsx").read("xl/_rels/workbook.xml.rels").decode()
    ok("template: workbook rels relative (Alteryx Overwrite Sheet stores sheets at xl//xl/... when they are absolute)",
       'Target="/' not in rels)
    shutil.copy(ROOT / "excel/templates/XoL_Results.xlsx", OUT / "XoL_Results.xlsx")
    run_all()
    book = OUT / "XoL_Results.xlsx"

    # data sheets vs staged tables and the spec
    wb = load_workbook(book, read_only=True)
    for sheet, cols in RS.SHEETS.items():
        ws = wb[sheet]
        head = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        ok(f"{sheet}: columns as results_spec", sorted(head) == sorted(cols))
        if sheet.startswith("out_") and (OUT / "stage" / f"res_{sheet}.yxdb").exists():
            fields, n = staged(f"res_{sheet}")
            ok(f"{sheet}: spec columns all produced by the stage", set(cols) <= set(fields))
            ok(f"{sheet}: {n} rows", ws.max_row - 1 == n)
    wb.close()

    # recalc a copy, with and without a loading override
    calc = OUT / "results_calc.xlsx"
    shutil.copy(book, calc)
    rc = recalc(calc)
    ok(f"LibreOffice recalculation: {rc.get('total_formulas')} formulas, 0 errors",
       rc.get("status") == "success" and rc.get("total_errors") == 0)
    premium = exhibit_checks(calc, OUT / "results", CTL)

    # loading override: +5% profit on PP
    ov = OUT / "results_override.xlsx"
    wb2 = load_workbook(book)
    wb2["Inputs"]["E30"] = 0.15
    wb2.save(ov)
    recalc(ov)
    w3 = load_workbook(ov, data_only=True)
    close("loading override flows into the PP 1 premium",
          [w3["Pricing"][f"{build_results.PC['Upfront annual premium']}{build_results.PROW[('PP', 1)]}"].value],
          [premium("PP", 1, load=0.05)["p"]])

    # scenario 1: rows of scenario 0 kept
    set_param("scenario_id", "1")
    run_all()
    sc = pd.read_excel(book, sheet_name="scn_summary")
    ok(f"scn_summary keeps scenario 0 and adds scenario 1 ({sorted(sc.scenario_id.unique())})",
       sorted(sc.scenario_id.unique()) == [0, 1] and (sc.scenario_id == 0).sum() == (sc.scenario_id == 1).sum())
    s1 = sc[(sc.scenario_id == 1) & (sc.struct_key == "1 PP 1")]
    s0 = sc[(sc.scenario_id == 0) & (sc.struct_key == "1 PP 1")]
    ok(f"scenario 1 (trend +2pp) raises the PP 1 expected loss ({s0.expected_loss_blended.iloc[0]:,.0f} -> "
       f"{s1.expected_loss_blended.iloc[0]:,.0f})", s1.expected_loss_blended.iloc[0] > s0.expected_loss_blended.iloc[0])
    log = pd.read_excel(book, sheet_name="out_run_log", dtype=str).set_index("item").value
    ok(f"run log: scenario '{log['scenario_name']}', version {log['tool_version']}, scenario_id {log['scenario_id']}",
       log["scenario_id"] == "1" and log["scenario_name"] == "Trend +2pp")
    for n, good in results:
        print(f"{'PASS' if good else 'FAIL'}  {n}")
    return 0 if all(g for _, g in results) else 1


if __name__ == "__main__":
    sys.exit(main())
